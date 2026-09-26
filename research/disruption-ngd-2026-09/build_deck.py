"""Build an editable, evidence-bound dark research presentation.

Run with /tmp/retraction-kdocs-20260916/venv/bin/python. Input is a JSON object
with status=verified, title, source_manifests=[{path, sha256}], and slides.
Each slide requires section, title, subtitle, source, and takeaway. It has
body (a list of paragraphs), chart, table, or columns (two lists of paragraphs).
Chart schema: type=scatter|bar|line, x_label, y_label, series=[{name, points:
[{x: number or category string, y: finite number, label: optional string}]}].
Null points must be omitted upstream and their exclusion disclosed in subtitle
or source. Source counts and denominators belong in the authored source text.
No publication or network operation occurs in this builder.
"""

import argparse
import hashlib
import json
import math
import unicodedata
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


PALETTE = {
    "background": "101C2D",
    "foreground": "F0F5FA",
    "muted": "B7C8D8",
    "rule": "41566F",
    "accent": "65D8EA",
    "series": ["65D8EA", "FFAC8B", "C5ABFA", "8CDEAE", "F2D175", "78A8F0"],
}
FONT = "Noto Sans CJK SC"
WIDTH = 16
HEIGHT = 9


def digest(path):
    value = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def contrast(first, second):
    def luminance(value):
        channels = [int(value[index:index + 2], 16) / 255 for index in (0, 2, 4)]
        linear = [channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4 for channel in channels]
        return sum(channel * weight for channel, weight in zip(linear, [0.2126, 0.7152, 0.0722]))
    levels = sorted([luminance(first), luminance(second)])
    return (levels[1] + 0.05) / (levels[0] + 0.05)


def validate_payload(payload):
    if payload.get("status") != "verified":
        raise ValueError("Only verified analysis payloads may produce a delivery deck")
    manifests = payload.get("source_manifests")
    if not manifests:
        raise ValueError("Source manifests are required")
    for manifest in manifests:
        if digest(manifest["path"]) != manifest["sha256"]:
            raise ValueError(f"Evidence hash mismatch: {manifest['path']}")
    slides = payload.get("slides", [])
    main_count = sum(not slide.get("appendix", False) for slide in slides)
    max_main_slides = payload.get("max_main_slides", 30)
    if isinstance(max_main_slides, bool) or not isinstance(max_main_slides, int) or not 30 <= max_main_slides <= 60:
        raise ValueError("Explicit max_main_slides must be an integer between 30 and 60")
    if not 24 <= main_count <= max_main_slides:
        raise ValueError(f"Expected 24–{max_main_slides} main slides, received {main_count}")
    for number, slide in enumerate(slides, 1):
        for field in ["section", "title", "subtitle", "source", "takeaway"]:
            if not isinstance(slide.get(field), str) or not slide[field].strip():
                raise ValueError(f"Slide {number}: missing {field}")
        if len(slide["title"]) > 48:
            raise ValueError(f"Slide {number}: title needs shortening")
        if not any(slide.get(field) for field in ["body", "columns", "chart", "table", "citation_demo"]):
            raise ValueError(f"Slide {number}: no substantive content")
        if slide.get("section_break") and (not slide.get("section_number") or len(slide.get("body", [])) != 2):
            raise ValueError("Chapter dividers require section_number and two transition paragraphs")
        if slide.get("chart"):
            chart = slide["chart"]
            if chart.get("type") not in ["scatter", "bar", "line"]:
                raise ValueError("Unsupported chart type")
            if not chart.get("x_label") or not chart.get("y_label"):
                raise ValueError("Chart units and axes must be supplied")
            if not chart.get("series"):
                raise ValueError("Empty chart")
            if len(chart["series"]) > len(PALETTE["series"]):
                raise ValueError("At most six series per readable figure")
            for series in chart["series"]:
                if not series.get("name") or not series.get("points"):
                    raise ValueError("Chart series cannot be empty")
                for point in series["points"]:
                    values = [point.get("y")]
                    if chart["type"] == "scatter":
                        values.append(point.get("x"))
                    if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) for value in values):
                        raise ValueError("Missing/nonfinite values cannot be silently zero-filled")
            if chart["type"] != "scatter":
                categories = [point["x"] for point in chart["series"][0]["points"]]
                if len(categories) > 18:
                    raise ValueError("Split categorical figures containing more than 18 categories")
                if any([point["x"] for point in series["points"]] != categories for series in chart["series"]):
                    raise ValueError("Category series must share the exact ordered categories")
        if slide.get("columns") and len(slide["columns"]) != 2:
            raise ValueError("Two-column layout requires exactly two columns")
    return {"main_slides": main_count, "total_slides": len(slides)}


def rgb(value):
    return RGBColor.from_string(value)


def text(slide, value, left, top, width, height, size=22, color="foreground", bold=False):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(PALETTE.get(color, color))
        paragraph.space_after = Pt(9)
        for run in paragraph.runs:
            east_asian = OxmlElement("a:ea")
            east_asian.set("typeface", FONT)
            run._r.get_or_add_rPr().append(east_asian)
    return shape


def rule(slide, left, top, width, color="rule"):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(0.015))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(PALETTE[color])
    shape.line.fill.background()


def body_text(slide, paragraphs, left, top, width, height, preferred_size=23):
    if not 21 <= preferred_size <= 25:
        raise ValueError("Body text must remain between 21 and 25 points")
    for size in range(preferred_size, 20, -1):
        line_count = sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(character) in "WF" else 0.58 for character in line) * size / (width * 72))) for paragraph in paragraphs for line in paragraph.split("\n"))
        estimated_height = line_count * size * 1.2 + (len(paragraphs) - 1) * 14
        if estimated_height <= height * 72:
            break
    else:
        raise ValueError("Body copy exceeds readable layout capacity; shorten or split the slide")
    shape = text(slide, "\n".join(paragraphs), left, top, width, height, size)
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(14 if index < len(shape.text_frame.paragraphs) - 1 else 0)
    return shape


def chapter_divider(slide, spec):
    text(slide, spec["section"], 0.8, 0.45, 14.2, 0.4, 15, "accent", True)
    text(slide, str(spec["section_number"]), 0.8, 1.35, 3.1, 1.6, 96, "accent", True)
    rule(slide, 4.0, 1.8, 11.2)
    text(slide, spec["title"], 4.0, 2.18, 11.1, 1.3, 36, bold=True)
    text(slide, spec["subtitle"], 4.0, 3.6, 11.1, 0.8, 21, "muted")
    body_text(slide, spec["body"], 4.0, 4.72, 11.1, 1.65, 23)


def diagram_node(slide, value, left, top, width, accent=False):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(0.65))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(PALETTE["background"])
    shape.line.color.rgb = rgb(PALETTE["accent"] if accent else PALETTE["rule"])
    shape.line.width = Pt(1.5)
    text(slide, value, left + 0.12, top + 0.13, width - 0.2, 0.4, 17, "accent" if accent else "foreground", True)


def diagram_arrow(slide, start_x, start_y, end_x, end_y, muted=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(start_x), Inches(start_y), Inches(end_x), Inches(end_y))
    connector.line.color.rgb = rgb(PALETTE["rule"] if muted else PALETTE["accent"])
    connector.line.width = Pt(1.5 if muted else 2.5)
    arrow = OxmlElement("a:tailEnd")
    arrow.set("type", "triangle")
    arrow.set("w", "med")
    arrow.set("len", "med")
    connector.line._get_or_add_ln().append(arrow)


def citation_diagram(slide):
    for index, (label, explanation, cite_focal, cite_reference) in enumerate([
        ("NF", "后来者只引用焦点论文", True, False),
        ("NB", "后来者同时引用二者", True, True),
        ("NR", "后来者只引用参考论文", False, True),
    ]):
        left = 0.8 + index * 4.92
        text(slide, label, left, 2.48, 4.45, 0.45, 28, "accent", True)
        text(slide, explanation, left, 3.05, 4.45, 0.45, 19)
        diagram_arrow(slide, left + 1.9, 4.125, left + 2.55, 4.125, True)
        if cite_focal:
            diagram_arrow(slide, left + 2.22, 5.5, left + 0.95, 4.45)
        if cite_reference:
            diagram_arrow(slide, left + 2.22, 5.5, left + 3.5, 4.45)
        diagram_node(slide, "焦点论文 F", left, 3.8, 1.9)
        diagram_node(slide, "参考论文 R", left + 2.55, 3.8, 1.9)
        diagram_node(slide, "后续论文 C", left + 1.27, 5.5, 1.9, True)
    text(slide, "箭头：引用者 → 被引用者。灰色 F → R 为既有参考关系；亮色箭头区分后来者的引用方式。", 0.8, 6.5, 14.3, 0.35, 15, "muted")


def chart_style(chart, spec):
    chart.has_title = False
    chart.has_legend = len(spec["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT
        chart.legend.font.size = Pt(13)
        chart.legend.font.color.rgb = rgb(PALETTE["foreground"])
    for axis in [chart.category_axis, chart.value_axis]:
        axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(13)
        axis.tick_labels.font.color.rgb = rgb(PALETTE["muted"])
        axis.format.line.color.rgb = rgb(PALETTE["rule"])
        axis.has_major_gridlines = False
    if spec["type"] != "scatter":
        if len(spec["series"][0]["points"]) >= 9:
            chart.category_axis.tick_labels.font.size = Pt(11)
        for tag in ["c:tickLblSkip", "c:tickMarkSkip"]:
            interval = OxmlElement(tag)
            interval.set("val", "1")
            chart.category_axis._element.insert_element_before(interval, "c:noMultiLvlLbl", "c:extLst")
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb(PALETTE["rule"])
    chart.value_axis.major_gridlines.format.line.width = Pt(0.4)
    for field, axis in [("x", chart.category_axis), ("y", chart.value_axis)]:
        if f"{field}_min" in spec:
            axis.minimum_scale = spec[f"{field}_min"]
        if f"{field}_max" in spec:
            axis.maximum_scale = spec[f"{field}_max"]
    chart.value_axis.tick_labels.number_format = spec.get("number_format", "0.00")
    for index, series in enumerate(chart.series):
        color = rgb(PALETTE["series"][index])
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.format.line.color.rgb = color
        if spec["type"] != "bar":
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 7
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.fill.background()
            if spec["type"] == "scatter":
                series.format.line.fill.background()
        labeled = [(point_index, point["label"]) for point_index, point in enumerate(spec["series"][index]["points"]) if point.get("label")]
        if labeled:
            labels = OxmlElement("c:dLbls")
            for point_index, label in labeled:
                element = OxmlElement("c:dLbl")
                identifier = OxmlElement("c:idx")
                identifier.set("val", str(point_index))
                element.append(identifier)
                value = OxmlElement("c:tx")
                rich = OxmlElement("c:rich")
                rich.append(OxmlElement("a:bodyPr"))
                rich.append(OxmlElement("a:lstStyle"))
                paragraph = OxmlElement("a:p")
                run = OxmlElement("a:r")
                properties = OxmlElement("a:rPr")
                properties.set("sz", "1100")
                fill = OxmlElement("a:solidFill")
                fill_color = OxmlElement("a:srgbClr")
                fill_color.set("val", PALETTE["foreground"])
                fill.append(fill_color)
                properties.append(fill)
                for tag in ["a:latin", "a:ea"]:
                    font = OxmlElement(tag)
                    font.set("typeface", FONT)
                    properties.append(font)
                run.append(properties)
                content = OxmlElement("a:t")
                content.text = label
                run.append(content)
                paragraph.append(run)
                rich.append(paragraph)
                value.append(rich)
                element.append(value)
                position = OxmlElement("c:dLblPos")
                position.set("val", "t")
                element.append(position)
                labels.append(element)
            for tag in ["c:showLegendKey", "c:showVal", "c:showCatName", "c:showSerName"]:
                flag = OxmlElement(tag)
                flag.set("val", "0")
                labels.append(flag)
            series._element.append(labels)
    chart_element = chart._chartSpace
    properties = chart_element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr")
    if properties is None:
        properties = OxmlElement("c:spPr")
        chart_element.append(properties)
    fill = OxmlElement("a:solidFill")
    fill_color = OxmlElement("a:srgbClr")
    fill_color.set("val", PALETTE["background"])
    fill.append(fill_color)
    properties.append(fill)
    border = OxmlElement("a:ln")
    border.append(OxmlElement("a:noFill"))
    properties.append(border)


def add_chart(slide, spec, display=None):
    if spec["type"] == "scatter":
        data = XyChartData()
        for source in spec["series"]:
            series = data.add_series(source["name"])
            for point in source["points"]:
                series.add_data_point(point["x"], point["y"])
        chart_type = XL_CHART_TYPE.XY_SCATTER
    else:
        data = CategoryChartData()
        data.categories = [str(point["x"]) for point in spec["series"][0]["points"]]
        for source in spec["series"]:
            data.add_series(source["name"], [point["y"] for point in source["points"]])
        chart_type = XL_CHART_TYPE.BAR_CLUSTERED if spec["type"] == "bar" else XL_CHART_TYPE.LINE_MARKERS
    chart = slide.shapes.add_chart(chart_type, Inches(0.75), Inches(2.4), Inches(14.4), Inches(4.05), data).chart
    chart_style(chart, spec)
    if display and 'number_format' in display:
        chart.value_axis.tick_labels.number_format = display['number_format']
    upper_label = spec["x_label"] if spec["type"] == "bar" else spec["y_label"]
    lower_label = spec["y_label"] if spec["type"] == "bar" else spec["x_label"]
    text(slide, upper_label, 0.8, 2.07, 14.2, 0.3, 14, "muted")
    text(slide, lower_label, 0.8, 6.5, 14.2, 0.32, 14, "muted")
    return chart


def add_table(slide, spec):
    columns, rows = spec["columns"], spec["rows"]
    if len(rows) > 9 or len(columns) > 6 or any(len(row) != len(columns) for row in rows):
        raise ValueError("Table requires at most nine rows/six columns and aligned values")
    table = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(0.8), Inches(2.25), Inches(14.4), Inches(4.35)).table
    widths = spec.get("column_widths")
    if widths is None:
        long_first = max((sum(2 if unicodedata.east_asian_width(character) in "WF" else 1 for character in str(row[0])) for row in rows), default=0) >= 26
        first_width = 4.9 if long_first and len(columns) > 2 else 14.4 / len(columns)
        widths = [first_width] + [(14.4 - first_width) / (len(columns) - 1)] * (len(columns) - 1) if len(columns) > 1 else [14.4]
    if len(widths) != len(columns) or any(width <= 0 for width in widths) or abs(sum(widths) - 14.4) > 0.01:
        raise ValueError("Table column widths must be positive and sum to 14.4 inches")
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row_index, values in enumerate([columns] + rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = "未估计" if value is None else str(value)
            cell.margin_top = Pt(1)
            cell.margin_bottom = Pt(1)
            cell.margin_left = Pt(5)
            cell.margin_right = Pt(5)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PALETTE["background"])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = FONT
                paragraph.font.size = Pt(16)
                paragraph.space_before = Pt(0)
                paragraph.space_after = Pt(0)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = rgb(PALETTE["accent"] if row_index == 0 else PALETTE["foreground"])


def build(payload, output):
    counts = validate_payload(payload)
    deck = Presentation()
    deck.slide_width = Inches(WIDTH)
    deck.slide_height = Inches(HEIGHT)
    deck.core_properties.title = payload["title"]
    deck.core_properties.subject = "撤稿学科 × 成熟窗口颠覆度 × 学科标签共现距离；探索性研究"
    deck.core_properties.author = "Scholarly Retraction Observatory"
    deck.core_properties.keywords = "OpenAlex, Retraction Watch, Disruption, NGD, exploratory"
    for number, spec in enumerate(payload["slides"], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(PALETTE["background"])
        if spec.get("section_break"):
            chapter_divider(slide, spec)
        else:
            text(slide, spec["section"], 0.8, 0.35, 14.2, 0.35, 13, "accent", True)
            text(slide, spec["title"], 0.8, 0.95, 14.3, 0.75, 31, bold=True)
            text(slide, spec["subtitle"], 0.8, 1.73, 14.3, 0.5, 16, "muted")
        if spec.get("section_break"):
            pass
        elif spec.get("citation_demo"):
            citation_diagram(slide)
        elif spec.get("chart"):
            add_chart(slide, spec["chart"], spec.get("chart_display"))
        elif spec.get("table"):
            add_table(slide, spec["table"])
        elif spec.get("columns"):
            for index, paragraphs in enumerate(spec["columns"]):
                left = 0.8 + index * 7.5
                rule(slide, left, 2.45, 6.7, "accent")
                body_text(slide, paragraphs, left, 2.72, 6.7, 3.9, spec.get("body_size", 23))
        else:
            body_text(slide, spec["body"], 0.8, 2.5, 14.3, 4.1, spec.get("body_size", 23))
        rule(slide, 0.8, 7.02, 14.4)
        text(slide, spec["takeaway"], 0.8, 7.18, 14.3, 0.67, 20, "accent", True)
        text(slide, spec["source"], 0.8, 8.07, 13.4, 0.59, 10, "muted")
        text(slide, f"{number:02d}", 14.45, 8.07, 0.65, 0.4, 15, "accent", True)
        slide.notes_slide.notes_text_frame.text = json.dumps({"slide": spec, "source_manifests": payload["source_manifests"], "interpretation": "Exploratory association, not causal evidence; missing values are not zero."}, ensure_ascii=False, indent=2)
    outside = []
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width or shape.top + shape.height > deck.slide_height:
                outside.append({"slide": number, "shape": shape.name})
    if outside:
        raise ValueError(f"Out-of-bounds shapes: {outside}")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    reopened = Presentation(output)
    validation = {**counts, "pptx": str(output.resolve()), "sha256": digest(output), "source_manifests": payload["source_manifests"], "native_chart_count": sum(shape.has_chart for slide in reopened.slides for shape in slide.shapes), "out_of_bounds": outside, "text_contrast": {name: contrast(PALETTE[name], PALETTE["background"]) for name in ["foreground", "muted", "accent"]}, "render_check": "required separately; bounds are not visual acceptance"}
    output.with_suffix(".validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    return validation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("payload", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = json.loads(args.payload.read_text(encoding="utf-8"))
    print(json.dumps(build(payload, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
