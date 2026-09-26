"""Render a built PPTX and audit text overflow; images still need human review."""

import argparse
import json
import subprocess
import tempfile
from pathlib import Path

import fitz
from pptx import Presentation


def render(source, output):
    source, output = Path(source).resolve(), Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="research-deck-lo-") as profile:
        subprocess.run(["libreoffice", f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(source)], check=True)
    pdf_path = output / source.with_suffix(".pdf").name
    document = fitz.open(pdf_path)
    presentation = Presentation(source)
    problems = []
    for number, page in enumerate(document, 1):
        page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25)).save(output / f"slide-{number:02d}.png")
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    left, top, right, bottom = span["bbox"]
                    if left < 0 or top < 0 or right > page.rect.width or bottom > page.rect.height:
                        problems.append({"slide": number, "issue": "text_outside_page", "text": span["text"], "bbox": span["bbox"]})
                    if 2.25 * 72 <= top < 7.18 * 72 and bottom > 7.02 * 72:
                        problems.append({"slide": number, "issue": "body_text_crosses_takeaway_rule", "text": span["text"], "bbox": span["bbox"]})
        for shape in presentation.slides[number - 1].shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if len(cell.text) < 12:
                            continue
                        for rectangle in page.search_for(cell.text):
                            if rectangle.y1 > 7.02 * 72 and rectangle.y0 < 7.9 * 72:
                                problems.append({"slide": number, "issue": "rendered_table_text_below_body", "text": cell.text, "bbox": list(rectangle)})
    result = {"pdf": str(pdf_path), "pages": len(document), "text_overflow": problems, "visual_review_required": True}
    (output / "render-audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = render(args.pptx, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["text_overflow"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
