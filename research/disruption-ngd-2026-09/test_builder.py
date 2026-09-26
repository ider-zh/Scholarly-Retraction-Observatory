import importlib.util
import tempfile
import unittest
from pathlib import Path

try:
    from pptx import Presentation
except ImportError as error:
    raise unittest.SkipTest("python-pptx is provided by the presentation virtual environment") from error


SPEC = importlib.util.spec_from_file_location("build_deck", Path(__file__).with_name("build_deck.py"))
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class BuilderTests(unittest.TestCase):
    def payload(self, count=24):
        source = Path(__file__)
        slide = dict(section="结构测试", title="不代表研究结果", subtitle="验证布局", source="单元测试", takeaway="不用于科学结论", body=["测试正文"])
        return dict(status="verified", title="TEST ONLY", source_manifests=[dict(path=str(source), sha256=BUILDER.digest(source))], slides=[dict(slide) for _ in range(count)])

    def test_explicit_expansion_cap(self):
        payload = self.payload(44)
        with self.assertRaises(ValueError):
            BUILDER.validate_payload(payload)
        payload["max_main_slides"] = 48
        self.assertEqual(BUILDER.validate_payload(payload)["main_slides"], 44)
        for invalid in [61, True, 29, "48"]:
            payload["max_main_slides"] = invalid
            with self.assertRaises(ValueError):
                BUILDER.validate_payload(payload)

    def test_editable_teaching_layouts(self):
        payload = self.payload()
        payload["slides"][0].update(section_break=True, section_number="01", body=["先解释问题", "再阅读证据"])
        payload["slides"][1].pop("body")
        payload["slides"][1]["citation_demo"] = True
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.pptx"
            validation = BUILDER.build(payload, path)
            self.assertEqual(validation["out_of_bounds"], [])
            deck = Presentation(path)
            self.assertEqual(len(deck.slides), 24)
            demo = deck.slides[1]
            self.assertFalse(any(shape.shape_type == 13 for shape in demo.shapes))
            self.assertEqual(sum(len(shape._element.xpath(".//a:tailEnd")) for shape in demo.shapes), 7)
            self.assertIn("引用者 → 被引用者", " ".join(shape.text for shape in demo.shapes if shape.has_text_frame))

    def test_evidence_and_missing_values_still_rejected(self):
        payload = self.payload()
        payload["slides"][0]["chart"] = dict(type="scatter", x_label="X", y_label="Y", series=[dict(name="test", points=[dict(x=1, y=None)])])
        with self.assertRaises(ValueError):
            BUILDER.validate_payload(payload)
        payload = self.payload()
        payload["source_manifests"][0]["sha256"] = "invalid"
        with self.assertRaises(ValueError):
            BUILDER.validate_payload(payload)

    def test_chart_display_precision_preserves_values(self):
        payload = self.payload()
        payload['slides'][0]['chart'] = dict(type='scatter', x_label='X', y_label='Y',
            series=[dict(name='test', points=[dict(x=0.005, y=0.004)])])
        payload['slides'][0]['chart_display'] = {'number_format': '0.000'}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.pptx'
            BUILDER.build(payload, path)
            deck = Presentation(path)
            chart = next(shape.chart for shape in deck.slides[0].shapes if shape.has_chart)
            self.assertEqual(chart.value_axis.tick_labels.number_format, '0.000')
            self.assertEqual(tuple(chart.series[0].values), (0.004,))


if __name__ == "__main__":
    unittest.main()
