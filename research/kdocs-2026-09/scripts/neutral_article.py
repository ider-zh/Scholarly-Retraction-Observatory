from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path('/tmp/retraction-nick-20260920')
source=(root/'中性综述_OpenAlex撤稿的学科分布与年度变化.md').read_text()
document=Document()
section=document.sections[0]
section.page_width=Cm(21)
section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(2)
section.left_margin=section.right_margin=Cm(2.2)
normal=document.styles['Normal']
normal.font.name='Noto Serif CJK SC'
normal.font.size=Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
normal.paragraph_format.line_spacing=1.22
normal.paragraph_format.space_after=Pt(6)
for line in source.splitlines():
    if not line or line=='---':continue
    if line.startswith('# '):
        paragraph=document.add_paragraph(line[2:])
        paragraph.runs[0].font.size=Pt(19)
        paragraph.runs[0].bold=True
        paragraph.paragraph_format.space_after=Pt(10)
    elif line.startswith('## '):
        paragraph=document.add_paragraph(line[3:])
        paragraph.paragraph_format.keep_with_next=True
        paragraph.paragraph_format.space_before=Pt(10)
        paragraph.runs[0].font.size=Pt(13)
        paragraph.runs[0].bold=True
    elif line.startswith('资料来源：'):
        document.add_paragraph('资料来源：')
        for name,url in re.findall(r'\[([^]]+)\]\(([^)]+)\)',line):
            paragraph=document.add_paragraph()
            paragraph.add_run(name+'：'+url).font.size=Pt(9)
    else:
        paragraph=document.add_paragraph(line)
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('第 ')
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
footer.add_run(' 页｜研究方法与结果综述')
document.core_properties.title='OpenAlex 撤稿的学科分布与年度变化：研究方法与结果综述'
document.save(root/'OpenAlex撤稿研究_方法与结果综述.docx')
print('source characters',len(source))
