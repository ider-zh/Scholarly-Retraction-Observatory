import json
from pathlib import Path

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation

OUT = Path('/tmp/retraction-kdocs-20260916')
presentation = Presentation(OUT / '撤稿集中在哪里_学科专题_2000-2025.pptx')
assert len(presentation.slides) == 30
for index, slide in enumerate(presentation.slides):
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (index, shape.name)
        assert shape.left+shape.width <= presentation.slide_width+100, (index, shape.name)
        assert shape.top+shape.height <= presentation.slide_height+100, (index, shape.name)
document = fitz.open(OUT / '撤稿集中在哪里_学科专题_2000-2025.pdf')
assert len(document) == 30
pages = []
for index, page in enumerate(document):
    assert len(page.get_text()) > 100, index
    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.25,1.25))
    path = OUT / f'page-{index+1:02d}.png'
    pixmap.save(path)
    pages.append(path)
for start in range(0,30,6):
    sheet = Image.new('RGB',(1440,3*427),'#e0e4e5')
    draw = ImageDraw.Draw(sheet)
    for offset, path in enumerate(pages[start:start+6]):
        image = Image.open(path).convert('RGB')
        image.thumbnail((710,400))
        left, top = (offset%2)*720, (offset//2)*427
        sheet.paste(image,(left,top+22))
        draw.text((left+10,top+3),str(start+offset+1),fill='black')
    sheet.save(OUT / f'contact-{start+1:02d}.jpg',quality=92)
print(json.dumps({'pages':len(document),'all_shapes_in_bounds':True,'rendered_all_pages':True,'contact_sheets':5}))
