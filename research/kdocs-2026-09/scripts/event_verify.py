import collections
import csv
import hashlib
import json
from pathlib import Path

import duckdb
from pptx import Presentation

BASE = Path('/tmp/retraction-kdocs-20260916')
OUT = BASE / 'event-year'
event = json.loads((OUT / 'analysis.json').read_text())
published = json.loads((BASE / 'published.json').read_text())
source = Path('/mnt/hg02/openalex-snapshot/analysis/broad-runs') / published['manifest']['scan_config_sha256']
papers = json.loads((source / 'rw_original.json').read_text())
matches = json.loads((source / 'rw_oa_match.json').read_text())
connection = duckdb.connect(config={'threads': 32, 'memory_limit': '128GB'})
connection.execute('CREATE TABLE links(id VARCHAR, date_text VARCHAR)')
connection.execute('BEGIN TRANSACTION')
connection.executemany('INSERT INTO links VALUES (?, ?)', [(matches[paper['id']]['selected_id'], paper['retracted']) for paper in papers if matches[paper['id']]['selected_id']])
connection.execute('COMMIT')
connection.read_parquet(str(source / 'shards/*/identifiers.parquet')).create_view('identities')
connection.read_parquet(str(source / 'shards/*/targets.parquet')).create_view('targets')
connection.execute("""CREATE TABLE selected AS
    WITH earliest AS (SELECT id, min(try_cast(date_text AS DATE)) AS retracted FROM links GROUP BY id)
    SELECT identities.id, year(retracted)::INTEGER AS year, targets.primary_topic.field.id AS field,
    targets.primary_topic.subfield.id AS subfield
    FROM earliest JOIN identities USING(id) LEFT JOIN targets USING(id)
    WHERE year(retracted) BETWEEN 2000 AND 2025
    AND identities.is_xpac IS FALSE AND identities.document_role IN ('original_supported','unresolved')
    AND identities.publication_year BETWEEN 1 AND 2026
    AND (identities.publication_date IS NULL OR identities.publication_date <= DATE '2026-06-26')""")
assert connection.execute('SELECT count(*),count(DISTINCT id) FROM selected').fetchone() == (58345,58345)
annual = dict(connection.execute('SELECT year,count(*) FROM selected GROUP BY year').fetchall())
assert [annual.get(year, 0) for year in event['summary']['years']] == event['summary']['annual_denominators']
independent = {}
for column, group in [('field','topics_field'),('subfield','topics_subfield')]:
    independent[group] = {(identifier, year): total for identifier, year, total in connection.execute(f"SELECT coalesce({column},'missing'),year,count(*) FROM selected GROUP BY 1,2").fetchall()}
concept_path = json.loads((source / 'concepts-complete.json').read_text())['directory']
connection.read_parquet(str(Path(concept_path) / '*.parquet')).create_view('concepts')
connection.execute('CREATE TABLE selected_concepts AS SELECT id,year,concepts FROM selected JOIN concepts USING(id)')
for level, group in [(0,'concepts_l0'),(1,'concepts_l1')]:
    pairs = connection.execute(f"""WITH expanded AS (SELECT DISTINCT id,year,tag.id AS label
        FROM selected_concepts,unnest(concepts) AS attached(tag) WHERE tag.level={level})
        SELECT label,year,count(*) FROM expanded GROUP BY 1,2
        UNION ALL SELECT 'missing',year,count(*) FROM selected_concepts
        WHERE NOT EXISTS (SELECT 1 FROM unnest(concepts) AS attached(tag) WHERE tag.level={level}) GROUP BY year""").fetchall()
    independent[group] = {(identifier, year): total for identifier, year, total in pairs}
checked = 0
with (OUT / 'all-subjects.csv').open(encoding='utf-8-sig') as handle:
    for row in csv.DictReader(handle):
        if row['retraction_year'] == '2000-2025':
            expected = sum(value for (identifier, year), value in independent[row['system']].items() if identifier == row['id'])
            denominator = 58345
        else:
            year = int(row['retraction_year'])
            expected = independent[row['system']].get((row['id'], year), 0)
            denominator = annual.get(year, 0)
        assert int(row['n']) == expected and int(row['N']) == denominator, row
        assert (not row['percent'] and denominator == 0) or abs(float(row['percent']) - expected / denominator * 100) < 1e-10
        checked += 1
before = Presentation(BASE / 'revision-audit/publication-reviewed.pptx')
after = Presentation(OUT / 'retraction-event-year.pptx')
assert len(after.slides) == 37
for original_index, original_slide in enumerate(before.slides):
    new_index = original_index if original_index < 24 else original_index + 7
    updated_slide = after.slides[new_index]
    old_images = [hashlib.sha256(shape.image.blob).hexdigest() for shape in original_slide.shapes if shape.shape_type == 13]
    new_images = [hashlib.sha256(shape.image.blob).hexdigest() for shape in updated_slide.shapes if shape.shape_type == 13]
    assert old_images == new_images
    if original_index not in [0,29]:
        original_text = [shape.text for shape in original_slide.shapes if shape.has_text_frame and not (shape.left > 13258800 and shape.top > 7315200)]
        updated_text = [shape.text for shape in updated_slide.shapes if shape.has_text_frame and not (shape.left > 13258800 and shape.top > 7315200)]
        assert original_text == updated_text, original_index
    original_tables = [[[cell.text for cell in row.cells] for row in shape.table.rows] for shape in original_slide.shapes if shape.has_table]
    updated_tables = [[[cell.text for cell in row.cells] for row in shape.table.rows] for shape in updated_slide.shapes if shape.has_table]
    assert original_tables == updated_tables
checks = {'independent_sql_aggregate_rows_checked': checked, 'all_old_chart_images_identical': True,
    'all_old_tables_identical': True, 'old_body_text_unchanged_except_cover_and_sources': True,
    'new_sample_unique_works': 58345, 'slides': 37}
(OUT / 'verification.json').write_text(json.dumps(checks, indent=2))
print(json.dumps(checks))
