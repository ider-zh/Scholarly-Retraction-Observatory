import os

os.environ['PPT_OUTPUT_DIR'] = '/tmp/retraction-nick-20260920'
helper = open('/tmp/retraction-kdocs-20260916/build_deck.py').read().split("item=slide('撤稿集中在哪里？'")[0]
exec(compile(helper,'existing-helpers','exec'))

import hashlib
import re
from docx import Document
from docx.shared import Pt as WordPt

SOURCE = pathlib.Path('/tmp/retraction-nick-20260920-v1')
DECK = Presentation(SOURCE/'学科撤稿研究_尼克简报_20260920.pptx')
original = list(DECK.slides)
data=json.loads((OUT/'analysis.json').read_text())
groups=data['groups']
base=np.array(data['summary']['annual_denominators'])
colors=[BLUE,RED,'#7C6B9B','#5C8861','#A7822F','#477D81']
TRANSLATIONS.update({'World Wide Web':'万维网','Physiology':'生理学','Biomedical Engineering':'生物医学工程',
    'Industrial and Manufacturing Engineering':'工业与制造工程','Control and Systems Engineering':'控制与系统工程',
    'Mechanical Engineering':'机械工程','Civil and Structural Engineering':'土木与结构工程',
    'Electrical and Electronic Engineering':'电气与电子工程'})

def update(shape,value):
    paragraph=shape.text_frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text=value
        for run in list(paragraph.runs)[1:]:run._r.getparent().remove(run._r)
    else:paragraph.text=value
    for extra in list(shape.text_frame.paragraphs)[1:]:extra._p.getparent().remove(extra._p)

def clear(item):
    for shape in list(item.shapes):shape._element.getparent().remove(shape._element)

clear(original[0])
text(original[0],'OpenAlex撤稿的学科分布与年度变化',.9,3.12,14.2,.9,36,True)
text(original[0],'基于 RW—OpenAlex 匹配样本的撤稿年份观察',.93,4.32,14.0,.7,23,False,MUTED)

clear(original[1])
text(original[1],'数据样本选择',.75,.95,14.5,.65,32,True)
text_columns(original[1],[
    ('OpenAlex：学科与文献','采用已验收的\n2026-06-26 快照。\n提供文献记录以及 Concepts、Topics 两套分类。'),
    ('RW：补充撤稿时间','采用 2026-09-10 快照。\n为按撤稿年份观察提供撤稿记录与日期。'),
    ('如何匹配','按 DOI / PMID 精确匹配，按不同 OA Work 去重。\n共同资格筛选后，2000—2025 年撤稿记录样本为 58,345 条。')
],top=2.55)
text(original[1],'正文使用匹配样本，不要求 OA 同时标记撤稿；附录仅使用 OA 发表年份样本。',.75,7.50,14.5,.45,15,False,MUTED)
line(original[1],.75,8.2,14.5,'#D1D8D8',.012)
text(original[1],'正文 / RW—OA 匹配样本 58,345 · OA 2026-06-26 / RW 2026-09-10',.75,8.38,13.6,.3,11,False,MUTED)
text(original[1],'02',14.65,8.32,.6,.35,15,True,BLUE)

selection=[]
evidence=list(csv.DictReader((OUT/'all-subjects.csv').open(encoding='utf-8-sig')))

def child_slide(parent,group):
    children=sorted([node for node in groups[group]['nodes'] if parent['id'] in node['parents']],key=lambda node:(-node['total'],node['id']))[:6]
    assert len(children)==6
    concepts=group=='concepts_l1'
    title_name=label(parent)
    item=slide(f'{title_name}：'+('关联二级概念的年度变化' if concepts else '所属 Subfield 的年度变化'),
        '正文 / RW—OPENALEX 匹配样本 / 2000—2025 撤稿记录年',
        f"先取{'一级学科' if concepts else 'Field'}全期数量前三，再在本组{'关联标签' if concepts else '子学科'}中固定选取数量前六；不是全局二级前六。",
        'Concepts 按分类树关联选取；每个 L1 仍计其全体文献，非与 L0 交集。多标签可重叠。' if concepts else
        '沿 primary_topic 的 Field → Subfield 路径选取；比例分母为全样本，不是所属 Field。')
    item.shapes[1].text_frame.paragraphs[0].font.size=Pt(27)
    figure,axes=plt.subplots(1,2,figsize=(14.55,5.1))
    figure.subplots_adjust(left=.06,right=.98,bottom=.15,top=.77,wspace=.17)
    for child,shade in zip(children,colors):
        annual=np.array(child['counts'])
        shares=np.divide(annual*100.,base,out=np.full(26,np.nan),where=base>0)
        axes[0].plot(YEARS,annual,color=shade,linewidth=1.9,label=label(child))
        axes[1].plot(YEARS,shares,color=shade,linewidth=1.9)
        small=(base>0)&(base<100)
        axes[1].scatter(YEARS[small],shares[small],s=28,edgecolors=shade,facecolors=BG,zorder=4)
    for axis,heading in zip(axes,['当年撤稿文献数（条）','年度撤稿文献中的学科占比（%）']):
        axis.set_title(heading,loc='left',fontsize=12)
        axis.set_xlim(2000,2025);axis.set_ylim(bottom=0)
        axis.set_xticks([2000,2005,2010,2015,2020,2025]);axis.set_xlabel('撤稿记录年')
        axis.grid(axis='y',color='#E4E8E6');axis.tick_params(length=0)
    axes[0].yaxis.set_major_formatter(ticker.StrMethodFormatter('{x:,.0f}'))
    handles,names=axes[0].get_legend_handles_labels()
    figure.legend(handles,names,loc='upper center',bbox_to_anchor=(.5,1.01),ncol=3,frameon=False,fontsize=11)
    figure.text(.06,.015,'空心点：年度总样本 N<100。占比下降不一定意味着数量减少；精确 n/N 见备注。',fontsize=10,color=MUTED)
    path=OUT/f"children-{group}-{parent['id'].split('/')[-1]}.png"
    figure.savefig(path,dpi=180);plt.close(figure)
    item.shapes.add_picture(str(path),Inches(.72),Inches(2.3),width=Inches(14.55),height=Inches(4.8))
    selected_ids={node['id'] for node in children}
    item.notes_slide.notes_text_frame.text=json.dumps({'group':group,'parent_id':parent['id'],'parent_label':parent['label'],
        'parent_total':parent['total'],'parent_selection':'top 3 by full-period count','child_selection':'top 6 within linked children',
        'nonintersection':concepts,'N_definition':'all eligible matched retracted Works in same event year, incl missing labels',
        'evidence':[row for row in evidence if row['system']==group and row['id'] in selected_ids]},ensure_ascii=False)
    selection.append({'group':group,'parent':parent,'children':children})
    return item

for parent_group,child_group in [('concepts_l0','concepts_l1'),('topics_field','topics_subfield')]:
    parents=groups[parent_group]['nodes'][:3]
    for parent in parents:child_slide(parent,child_group)

for shape in original[9].shapes:
    if shape.has_text_frame and shape.text.startswith('Concepts L1 法学标签'):
        update(shape,'二级分析先选高数量主学科，再展开其关联标签或所属子学科。\nConcepts 是多标签关联，Topics 是主主题路径。\n两者不能按同名分类直接等同。')
for shape in original[10].shapes:
    if shape.has_text_frame and shape.text.startswith('Concepts 延续前期研究'):
        update(shape,'Concepts 一级前三为计算机科学、医学、生物学。\nTopics Field 前三为医学、生物化学／遗传与分子生物学、工程学。\n分别在这三组中展开二级分析。')

identifiers=list(DECK.slides._sldIdLst)
order=list(range(5))+[15,16,17,6,7,18,19,20,9,10,11,12,13,14]
for index,identifier in enumerate(identifiers):
    DECK.slides._sldIdLst.remove(identifier)
    if index not in order:DECK.part.drop_rel(identifier.rId)
for index in order:DECK.slides._sldIdLst.append(identifiers[index])
distribution_pages={4,9}
trend_pages={5,6,7,8,10,11,12,13}
appendix_pages={16,17,18,19}
for number,item in enumerate(DECK.slides,1):
    for shape in item.shapes:
        if shape.has_text_frame and shape.left>Inches(14.5) and shape.top>Inches(8):update(shape,f'{number:02d}')
        if shape.shape_type==13 and number in distribution_pages|trend_pages|appendix_pages:
            shape.height=Inches(4.8)
    if number in distribution_pages:
        formula='数量 n＝全期该学科匹配撤稿文献数；全期学科占比＝n ÷ 58,345 × 100%。'
    elif number in trend_pages:
        formula='年度学科占比＝当年该学科匹配撤稿文献数 n ÷ 当年全部合格匹配撤稿文献数 N × 100%。'
    elif number in appendix_pages:
        formula='OA 学科内撤稿标记比例＝同学科撤稿标记文献数 n ÷ 同学科全部合格 OA 发文数 N × 100%。'
        if number in {17,19}:formula='OA 年度学科内比例＝同发表年、同学科撤稿标记文献数 n ÷ 同发表年、同学科全部合格发文数 N × 100%。'
    else:formula=None
    if formula:text(item,formula,.75,7.18,14.5,.45,14,False,BLUE)
    if 6<=number<=8 or 11<=number<=13:
        for shape in item.shapes:
            if shape.has_text_frame and shape.text.startswith('SCHOLARLY RETRACTION OBSERVATORY'):
                update(shape,'正文 / RW—OA 匹配样本 58,345 · OA 2026-06-26 / RW 2026-09-10')
    if item.has_notes_slide:
        item.notes_slide.notes_text_frame.text=re.sub(r'"page":\s*\d+',f'"page": {number}',item.notes_slide.notes_text_frame.text)
assert len(DECK.slides)==19
DECK.save(OUT/'学科撤稿研究_尼克简报_20260920.pptx')
(OUT/'parent-child-selection.json').write_text(json.dumps(selection,ensure_ascii=False,indent=2))

document=Document(SOURCE/'学科撤稿研究_综述_20260920.docx')

def replace_paragraph(paragraph,value):
    paragraph.text=value

for paragraph in document.paragraphs:
    value=paragraph.text
    value=value.replace('第 12—15 页','第 16—19 页')
    value=value.replace('前六按全期数量固定选取；','一级学科／Field 趋势取全期前六；二级分析先取一级学科／Field 前三，再在各组中选择二级前六；')
    if value=='2. 二级概念：数量与年度变化':value='2. 在前三个一级学科中展开二级概念'
    if value.startswith('法学标签全期关联 13,214'):
        value='二级分析不再选择全局 L1 前六，而是先选计算机科学、医学、生物学三个高数量一级学科，再按分类树关联选取各组 L1 数量前六，分别呈现年度数量和占比。下表列出各组前三项，完整六项趋势见 PPT 第 6—8 页。'
    if value.startswith('局部交叉检查显示，2011 年带 Law'):
        value='这里的父子关系用于组织与导航；L1 数量仍是该标签全体文献，不是同时带有父 L0 的交集。同一标签可出现在多个主学科页面，不能重复相加。例如癌症研究关联医学与生物学；所有年度占比仍以当年全部合格匹配撤稿文献为分母。'
    if value=='2. Subfield 与年度变化':value='2. 在前三个 Field 中展开 Subfield'
    if value.startswith('分子生物学、癌症研究与人工智能位列'):
        value='先选医学、生物化学／遗传与分子生物学、工程学三个 Field，再沿 primary_topic 路径选择各自数量前六的 Subfield。下表列出各组三项，完整六项趋势见 PPT 第 11—13 页。占比分母仍是全部合格匹配撤稿文献，不改为所属 Field。'
    if value.startswith('对 2011 年 Law 标签峰值的局部检验'):
        value='二级分析采用先主类、后子类的组织方式，但 Concepts 与 Topics 的父子含义不同。Concepts 按分类树关联组织 L1，计数不与父标签取交集；Topics Subfield 则属于其主主题路径中的 Field。因而各组排名用于描述不同分类规则下的分布，不能将两套子类视为一一对应，也不能把 Concepts 子项相加解释成父项。'
    if value.startswith('对研究需求的回答是：'):
        value='对研究需求的回答是：Concepts 一级学科前三为计算机科学、医学、生物学；Topics Field 前三为医学、生物化学／遗传与分子生物学、工程学。在这六组中分别展开关联 L1 或所属 Subfield 的年度变化，更便于观察高数量主类涉及的具体研究方向，而非依赖全局二级排名。排名必须与分类定义一并解释，不能作为“高风险学科榜”。'
    value=value.replace('下表列出各组前三项','表中列出各组前三项').replace('下表列出各组三项','表中列出各组三项')
    if value!=paragraph.text:replace_paragraph(paragraph,value)

for table_index,group in [(2,'concepts_l1'),(4,'topics_subfield')]:
    grid=document.tables[table_index]
    for row in list(grid.rows):grid._tbl.remove(row._tr)
    rows=[['主学科／Field · 二级标签','记录数','全期占比']]
    for entry in selection:
        if entry['group']!=group:continue
        parent_label=label(entry['parent'])
        if parent_label=='生物化学、遗传与分子生物学':parent_label='生化／遗传／分子生物学'
        for child in entry['children'][:3]:rows.append([parent_label+' · '+label(child),f"{child['total']:,}",f"{child['total']/58345*100:.2f}%"])
    for values in rows:
        for cell,value in zip(grid.add_row().cells,values):
            cell.text=value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=WordPt(2)
                for run in paragraph.runs:run.font.size=WordPt(9)
    for column_index, width in enumerate([4.5, .9, .9]):
        grid.columns[column_index].width=int(914400*width)
        for row in grid.rows:
            row.cells[column_index].width=int(914400*width)
document.save(OUT/'学科撤稿研究_综述_20260920.docx')

for entry in selection:
    group=entry['group'];parent=entry['parent'];children=entry['children']
    assert all(parent['id'] in child['parents'] for child in children)
    expected=sorted([node for node in groups[group]['nodes'] if parent['id'] in node['parents']],key=lambda node:(-node['total'],node['id']))[:6]
    assert children==expected
    if group=='topics_subfield':
        all_children=[node for node in groups[group]['nodes'] if parent['id'] in node['parents']]
        assert [sum(child['counts'][index] for child in all_children) for index in range(26)]==parent['counts']
    print(parent['label'],[(child['label'],child['total']) for child in children])
print('19 slides; six parent-based drilldowns; underlying aggregate data unchanged.')
