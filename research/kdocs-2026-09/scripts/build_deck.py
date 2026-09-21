import csv
import json
import os
import pathlib
import textwrap

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager, ticker
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

BASE = pathlib.Path('/tmp/retraction-kdocs-20260916')
OUT = pathlib.Path(os.environ.get('PPT_OUTPUT_DIR', str(BASE)))
OUT.mkdir(exist_ok=True)
DATA = json.loads((BASE / 'published.json').read_text())
ANALYSIS = json.loads((BASE / 'analysis.json').read_text())
TAX = {item['id']: item for item in DATA['fields']['discipline_explorer']['taxonomies']}
YEARS = np.arange(2000, 2026)
BLUE, RED, INK, MUTED, BG = '#176B87', '#B45C42', '#182B35', '#61717A', '#FAFAF7'
COLORS = [BLUE, RED, '#7C6B9B', '#5C8861', '#A7822F']
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'font.size': 12, 'axes.unicode_minus': False,
    'axes.spines.top': False, 'axes.spines.right': False, 'axes.spines.left': False,
    'axes.edgecolor': '#CAD1D3', 'text.color': INK, 'axes.labelcolor': MUTED,
    'xtick.color': MUTED, 'ytick.color': INK, 'figure.facecolor': BG, 'axes.facecolor': BG,
    'savefig.facecolor': BG})

TRANSLATIONS = {
    'Computer science':'计算机科学', 'Computer Science':'计算机科学', 'Medicine':'医学', 'Biology':'生物学',
    'Chemistry':'化学', 'Engineering':'工程学', 'Physics':'物理学', 'Geology':'地质学', 'Philosophy':'哲学',
    'Art':'艺术', 'Sociology':'社会学', 'Business':'商业', 'Psychology':'心理学', 'Economics':'经济学',
    'Political science':'政治学', 'Materials science':'材料科学', 'Geography':'地理学', 'Mathematics':'数学',
    'Environmental science':'环境科学', 'History':'历史学', 'Artificial intelligence':'人工智能',
    'Artificial Intelligence':'人工智能', 'Operating system':'操作系统', 'Machine learning':'机器学习',
    'Algorithm':'算法', 'Programming language':'编程语言', 'Data mining':'数据挖掘',
    'Internal medicine':'内科学', 'Cancer research':'癌症研究', 'Pathology':'病理学', 'Surgery':'外科学',
    'Immunology':'免疫学', 'Endocrinology':'内分泌学', 'Biochemistry':'生物化学', 'Genetics':'遗传学',
    'Cell biology':'细胞生物学', 'Organic chemistry':'有机化学', 'Chromatography':'色谱分析',
    'Physical chemistry':'物理化学', 'Food science':'食品科学', 'Nuclear chemistry':'核化学',
    'Telecommunications':'电信', 'Mechanical engineering':'机械工程', 'Electrical engineering':'电气工程',
    'Chemical engineering':'化学工程', 'Operations research':'运筹学', 'Structural engineering':'结构工程',
    'Agricultural and Biological Sciences':'农业与生物科学', 'Arts and Humanities':'艺术与人文学科',
    'Biochemistry, Genetics and Molecular Biology':'生物化学、遗传与分子生物学',
    'Business, Management and Accounting':'商业、管理与会计', 'Chemical Engineering':'化学工程',
    'Decision Sciences':'决策科学', 'Earth and Planetary Sciences':'地球与行星科学',
    'Economics, Econometrics and Finance':'经济、计量经济与金融', 'Energy':'能源',
    'Environmental Science':'环境科学', 'Immunology and Microbiology':'免疫与微生物学',
    'Materials Science':'材料科学', 'Neuroscience':'神经科学', 'Nursing':'护理学',
    'Pharmacology, Toxicology and Pharmaceutics':'药理、毒理与药剂学', 'Physics and Astronomy':'物理与天文学',
    'Social Sciences':'社会科学', 'Veterinary':'兽医学', 'Dentistry':'牙科学', 'Health Professions':'卫生专业',
    'Pulmonary and Respiratory Medicine':'肺与呼吸医学', 'Oncology':'肿瘤学', 'Epidemiology':'流行病学',
    'Radiology, Nuclear Medicine and Imaging':'放射、核医学与影像', 'Cardiology and Cardiovascular Medicine':'心脏与心血管医学',
    'Computer Networks and Communications':'计算机网络与通信', 'Information Systems':'信息系统',
    'Computer Vision and Pattern Recognition':'计算机视觉与模式识别', 'Signal Processing':'信号处理',
    'Computer Graphics and Computer-Aided Design':'计算机图形与辅助设计', 'Molecular Biology':'分子生物学',
    'Cancer Research':'癌症研究', 'Cell Biology':'细胞生物学', 'Molecular Medicine':'分子医学',
    'Clinical Biochemistry':'临床生物化学',
}

def label(node):
    value = node['label'] if isinstance(node, dict) else node
    return TRANSLATIONS.get(value, value)

def count(node, population='A1'):
    return sum(node['counts'][population][1:27])

def denominator(node):
    return sum(node['denominator'][1:27])

def proportion(node, population='A1'):
    return count(node, population) / denominator(node) * 100 if denominator(node) else None

def roots(taxonomy, level):
    return sorted([node for node in TAX[taxonomy]['nodes'] if node['level'] == level and not node.get('missing') and not node.get('navigation_only')], key=lambda node: -count(node))

CONCEPTS = roots('concepts', 0)
FIELDS = roots('topics', 1)
assert len(CONCEPTS) == 19 and len(FIELDS) == 26
TOP_CONCEPTS, TOP_FIELDS = CONCEPTS[:5], FIELDS[:5]
EVIDENCE = []
SLIDES = []
AXIS_AUDIT = []
DECK = Presentation()
DECK.slide_width, DECK.slide_height = Inches(16), Inches(9)
DECK.core_properties.title = '撤稿集中在哪里？从数量走向学科内比例'
DECK.core_properties.subject = 'OpenAlex 与匹配 Retraction Watch 的学科撤稿观察，2000–2025 发表队列'
DECK.core_properties.author = 'Scholarly Retraction Observatory'

def color(hex_value):
    return RGBColor.from_string(hex_value.lstrip('#'))

def text(slide, value, left, top, width, height, size=20, bold=False, fill=INK):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(value.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = 'Noto Sans CJK SC'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
        paragraph.space_after = Pt(10)
    return shape

def line(slide, left, top, width, fill=BLUE, height=0.025):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()

def slide(title, eyebrow, subtitle='', caption='', notes=''):
    item = DECK.slides.add_slide(DECK.slide_layouts[6])
    item.background.fill.solid()
    item.background.fill.fore_color.rgb = color(BG)
    text(item, eyebrow, .7, .35, 14.5, .35, 12, True, BLUE)
    text(item, title, .7, .92, 14.6, .62, 30, True)
    text(item, subtitle, .72, 1.68, 14.5, .57, 15, False, MUTED)
    line(item, .72, 8.20, 14.55, '#D1D8D8', .012)
    text(item, caption, .72, 7.65, 14.6, .50, 11, False, MUTED)
    text(item, 'SCHOLARLY RETRACTION OBSERVATORY  /  学科专题  ·  OA 快照 2026-06-26', .72, 8.38, 13.7, .25, 10, False, MUTED)
    text(item, f'{len(DECK.slides):02d}', 14.65, 8.32, .6, .35, 15, True, BLUE)
    item.notes_slide.notes_text_frame.text = notes
    SLIDES.append({'page': len(DECK.slides), 'title': title, 'source': eyebrow, 'caption': caption})
    return item

def note_evidence(item, nodes, population, taxonomy):
    rows = []
    for node in nodes:
        row = {'page': len(DECK.slides), 'taxonomy': taxonomy, 'id': node['id'], 'label': node['label'],
               'population': population, 'year': '2000–2025', 'n': count(node, population), 'N': denominator(node),
               'percent': proportion(node, population)}
        rows.append(row)
        EVIDENCE.append(row)
        for index, year in enumerate(YEARS, 1):
            numerator, base = node['counts'][population][index], node['denominator'][index]
            annual_row = dict(row, year=int(year), n=numerator, N=base, percent=numerator / base * 100 if base else None)
            EVIDENCE.append(annual_row)
            rows.append(annual_row)
    item.notes_slide.notes_text_frame.text += '\n数据证据（原始整数；比例未舍入；年度小基数不连线）：\n' + json.dumps(rows, ensure_ascii=False)

def axes_style(axis, horizontal=False):
    axis.grid(axis='x' if horizontal else 'y', color='#E4E8E6', linewidth=.7)
    axis.set_axisbelow(True)
    axis.tick_params(length=0, pad=6)
    if horizontal:
        axis.invert_yaxis()
    else:
        axis.set_xlim(2000, 2025)
        axis.set_xticks([2000, 2005, 2010, 2015, 2020, 2025])
        axis.set_xlabel('原文发表年', fontsize=11)
    axis.set_ylim(bottom=0) if not horizontal else None

def save_figure(figure, suffix):
    path = OUT / f'figure-{len(DECK.slides):02d}-{suffix}.png'
    figure.savefig(path, dpi=170)
    plt.close(figure)
    return path

def embed(item, path, left=.72, top=2.3, width=14.55, height=5.2):
    item.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))

def bar_pair(axes, nodes, population='A1', compact=False):
    yvalues = np.arange(len(nodes))
    values = [count(node, population) for node in nodes]
    rates = [proportion(node, population) for node in nodes]
    hue = BLUE if population == 'A1' else RED
    axes[0].barh(yvalues, values, color=hue, height=.58)
    axes[0].set_yticks(yvalues, [label(node) for node in nodes], fontsize=10 if len(nodes)>20 else 12)
    axes[0].set_title('撤稿记录数量  n（条）', loc='left', fontsize=13, pad=12)
    axes[0].xaxis.set_major_formatter(ticker.FuncFormatter(lambda value, pos:f'{value:,.0f}'))
    axes[0].set_xlim(0, max(values) * 1.30)
    for position, value in enumerate(values):
        axes[0].text(value + max(values)*.015, position, f'{value:,}', va='center', fontsize=9 if compact else 10)
    valid_rates = []
    for position, node in enumerate(nodes):
        rate = rates[position]
        if count(node, population)>=20 and denominator(node)>=1000 and rate is not None:
            axes[1].barh(position, rate, color=hue, alpha=.75, height=.58)
            axes[1].text(rate, position, f'  {rate:.3f}%', va='center', fontsize=9 if compact else 10)
            valid_rates.append(rate)
        else:
            axes[1].text(0, position, '小基数，不排名', va='center', fontsize=9, color=MUTED)
    axes[1].set_yticks(yvalues, ['']*len(nodes))
    axes[1].set_title('学科内撤稿比例  n / N（%）', loc='left', fontsize=13, pad=12)
    axes[1].set_xlim(0, max(valid_rates or [1]) * 1.35)
    for axis in axes:
        axes_style(axis, True)

def trends(axes, nodes, population='A1', legend=True, comparable=False):
    for index, node in enumerate(nodes):
        counts = np.array(node['counts'][population][1:27])
        bases = np.array(node['denominator'][1:27])
        ratios = np.divide(counts*100, bases, out=np.full(26, np.nan), where=bases>0)
        ratios[(counts<20)|(bases<1000)] = np.nan
        axes[0].plot(YEARS, counts, color=COLORS[index], linewidth=2, label=label(node))
        axes[1].plot(YEARS, ratios, color=COLORS[index], linewidth=2, label=label(node))
    axes[0].set_title('发表年队列：已观测撤稿数量（条）', loc='left', fontsize=13)
    axes[1].set_title('发表年队列：学科内撤稿比例（%）', loc='left', fontsize=13)
    axes[0].yaxis.set_major_formatter(ticker.StrMethodFormatter('{x:,.0f}'))
    for axis in axes:
        axes_style(axis)
        axis.tick_params(labelsize=10)
    if legend:
        axes[0].legend(loc='upper left', frameon=False, fontsize=10)
    if comparable:
        all_counts, all_ratios = [], []
        for node in nodes:
            for source in ['A1', 'C_D']:
                for numerator, base in zip(node['counts'][source][1:27], node['denominator'][1:27]):
                    all_counts.append(numerator)
                    if numerator >= 20 and base >= 1000:
                        all_ratios.append(numerator / base * 100)
        for axis, values in zip(axes, [all_counts, all_ratios]):
            ticks = ticker.MaxNLocator(nbins=5).tick_values(0, max(values) * 1.05)
            axis.set_ylim(0, ticks[-1])
            axis.set_yticks(ticks)
        AXIS_AUDIT.append({'page':len(DECK.slides), 'population':population,
            'nodes':[node['id'] for node in nodes],
            'limits':[list(axis.get_ylim()) for axis in axes],
            'ticks':[list(axis.get_yticks()) for axis in axes]})

COMMON = '2000–2025 发表队列；观察截至 2026-06-26。N 为同年、同分类、同宽口径的 OA 全部合格发文；并非不端发生率。'
SMALL = '比例仅绘制 n≥20 且 N≥1,000 的单元；缺口不等于零。近期队列随访较短，不能把下降解读为改善。'
CONCEPT_NOTE = 'Concepts 多标签可重叠；L1 数值为该标签全体记录，不是与父标签取交集。旧树只用于导航，子项不可相加推算父项。'

def distribution(title, eyebrow, nodes, population, taxonomy, subtitle='', caption=COMMON):
    item = slide(title, eyebrow, subtitle, caption)
    figure, axes = plt.subplots(1, 2, figsize=(14.55, 5.2), gridspec_kw={'width_ratios':[1.12,1]})
    figure.subplots_adjust(left=.19 if taxonomy=='topics' else .105, right=.96, bottom=.09, top=.92, wspace=.16)
    bar_pair(axes, nodes, population, compact=len(nodes)>20)
    embed(item, save_figure(figure, 'distribution'))
    note_evidence(item, nodes, population, taxonomy)
    return item

def trend_slide(title, eyebrow, nodes, population, taxonomy):
    subtitle = '同一组学科，同步观察数量与同学科、同发表年份发文中的比例。'
    if taxonomy == 'topics':
        subtitle += ' OA / RW 对应图采用相同纵轴刻度。'
    item = slide(title, eyebrow, subtitle, SMALL)
    figure, axes = plt.subplots(1, 2, figsize=(14.55, 5.2))
    figure.subplots_adjust(left=.065, right=.98, bottom=.12, top=.90, wspace=.20)
    trends(axes, nodes, population, comparable=taxonomy=='topics')
    embed(item, save_figure(figure, 'trend'))
    note_evidence(item, nodes, population, taxonomy)

def detail_slide(root, taxonomy, population):
    children = sorted([node for node in TAX[taxonomy]['nodes'] if root['id'] in node['parents'] and not node.get('missing')], key=lambda node:-count(node))[:6]
    source = 'OPENALEX' if population=='A1' else 'RW × OPENALEX 匹配样本'
    scope = 'CONCEPTS  /  主学科关联标签（非交集）' if taxonomy=='concepts' else 'TOPICS FIELD → SUBFIELD'
    title = f'{label(root)}：与该主学科关联的 L1 标签分析' if taxonomy=='concepts' else f'{label(root)}：哪些细分方向贡献更多记录？'
    subtitle = '按旧分类树关联选取 L1 标签；展示 OA 数量前 6，趋势追踪前 3。每项 n/N 均为该 L1 标签全体文献口径。' if taxonomy=='concepts' else f'展示 OA 数量前 6 个子学科；趋势追踪前 3，OA / RW 对应趋势图纵轴统一。父级 n={count(root,population):,}，N={denominator(root):,}。'
    item = slide(title, f'{source}  /  {scope}', subtitle,
        CONCEPT_NOTE if taxonomy=='concepts' else 'OA / RW 使用相同子学科清单，不另选各自 Top N。比例趋势仅连接 n≥20、N≥1,000 单元；均为发表年队列。')
    figure, axes = plt.subplots(2, 2, figsize=(14.55, 5.3), gridspec_kw={'height_ratios':[1,1]})
    figure.subplots_adjust(left=.17, right=.97, bottom=.09, top=.94, wspace=.22, hspace=.60)
    bar_pair(axes[0], children, population)
    trends(axes[1], children[:3], population, comparable=taxonomy=='topics')
    embed(item, save_figure(figure, 'detail'), height=5.3)
    note_evidence(item, [root]+children, population, taxonomy)

def text_columns(item, columns, top=2.7):
    width = 14.5/len(columns)
    for index, (heading, body) in enumerate(columns):
        left = .75+index*width
        line(item,left,top,width-.45)
        text(item,heading,left,top+.30,width-.5,.6,24,True)
        text(item,body,left,top+1.1,width-.55,3.7,20)

def table(item, headers, rows, left=.8, top=2.45, width=14.4, height=4.9, font_size=14):
    shape = item.shapes.add_table(len(rows)+1,len(headers), Inches(left), Inches(top), Inches(width), Inches(height))
    grid = shape.table
    for row_index, values in enumerate([headers]+rows):
        for column_index,value in enumerate(values):
            cell=grid.cell(row_index,column_index)
            cell.text=str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb=color('#E5EFEE' if row_index==0 else BG)
            cell.margin_top=Pt(5)
            cell.margin_bottom=Pt(4)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name='Noto Sans CJK SC'
                paragraph.font.size=Pt(font_size)
                paragraph.font.bold=row_index==0
                paragraph.font.color.rgb=color(INK)
    return grid

item=slide('撤稿集中在哪里？', '学科撤稿观察  /  RESEARCH BRIEF',
    '从数量走向学科内比例：OpenAlex 与匹配 Retraction Watch 的两条观察线',
    '分析单位为数据库合格 Work 记录；不是投稿退稿率，也不是学术不端发生率。')
text(item,'数量多的学科，\n比例也更高吗？',.8,2.8,10.5,2,46,True)
text(item,'2000—2025 发表队列\n共同观察截止：2026 年 6 月 26 日',.85,5.65,11,1.2,24,False,MUTED)
text(item,'77,414',12.0,3.0,3.1,.8,38,True,BLUE)
text(item,'OA 撤稿标记候选',12.0,3.9,3.1,.6,16)
text(item,'59,028',12.0,5,3.1,.8,38,True,RED)
text(item,'匹配 RW 合格文献',12.0,5.9,3.1,.6,16)

medicine=next(node for node in FIELDS if node['label']=='Medicine')
biochem=next(node for node in FIELDS if node['label']=='Biochemistry, Genetics and Molecular Biology')
item=slide('换一个分母，学科画像就会变化', '研究综述  /  先看三个结论', caption=COMMON)
text_columns(item,[('数量不等于比例',f'Topics 医学数量为 {count(medicine):,} 条，比例 {proportion(medicine):.3f}%。\n生化、遗传与分子生物学为 {count(biochem):,} 条，比例 {proportion(biochem):.3f}%。'),
    ('分类方式改变画像',f'Concepts 计算机科学关联 {count(TOP_CONCEPTS[0]):,} 条。\nTopics 主 Field 计算机科学为 {count(FIELDS[1]):,} 条。\n前者可多标签，后者每篇只归一个主类。'),
    ('匹配后才可比',f'RW 主分析使用 {ANALYSIS["rw_eligible_distinct_window"]:,} 条匹配记录。\n沿用同一 OA Field / Subfield 与发文分母。\n两库重叠且来源相关，不能相加。')])
note_evidence(item,[medicine,biochem,FIELDS[1]],'A1','topics')

item=slide('我们没有只保留 article，也没有使用整个 OA 不加筛选', '数据范围  /  WORK 是学术文献记录',
    'OpenAlex 包含论文、综述、书籍章节、书籍、数据集等；本研究采用主体库 core 的全类型宽口径。',
    '仅排除证据明确的独立通知及日期不合格项；标题、type 本身不作为排除依据。保留身份不确定项，不声称均为已确认原文。')
selected_types=ANALYSIS['types'][:7]
remaining_types = ANALYSIS['types'][7:]
type_rows = [[kind,f'{base:,}',f'{number:,}'] for kind,base,number in selected_types]
type_rows.append(['其余 17 类（合计）', f'{sum(row[1] for row in remaining_types):,}', f'{sum(row[2] for row in remaining_types):,}'])
type_rows.append(['全部类型合计',f'{ANALYSIS["denominator_window"]:,}',f'{ANALYSIS["oa_window"]:,}'])
assert sum(row[1] for row in ANALYSIS['types']) == ANALYSIS['denominator_window']
assert sum(row[2] for row in ANALYSIS['types']) == ANALYSIS['oa_window']
table(item,['文献类型（按发文量节选）','全部合格发文 N','OA 撤稿候选 n'],type_rows,width=9.1,font_size=14)
text(item,'220,305,891',10.6,2.7,4.5,.8,34,True,BLUE)
text(item,'共同发文分母\n2000–2025 年发表\n\nOA 快照：2026-06-26\nRW 快照：2026-09-10\nRW 事件截断：2026-06-26',10.6,3.65,4.5,3.2,19)
item.notes_slide.notes_text_frame.text='表内为发文数量前七类型，不是全部类型。全类型计数：'+json.dumps(ANALYSIS['types'],ensure_ascii=False)

item=slide('同学科、同发表年，分子与分母成对出现', '方法  /  比例到底除以什么', caption=SMALL)
text_columns(item,[('OA 观察线','n：学科内带撤稿标记的合格 Work。\nN：同学科全部合格 OA 发文。\n比例 = n / N × 100%。'),
    ('RW 观察线','n：RW 已匹配 OA、通过共同范围与事件截止的不同 Work。\nN：同一 OA 学科发文分母。\n不是除以 RW 撤稿总数。'),
    ('时间轴的含义','按原文发表年，不是撤稿发生年。\n每个年份都有自己的 n 和 N。\n近期发文随访较短；比例不是终身风险。')])

distribution('Concepts：计算机科学数量最多，不意味着比例最高', 'OPENALEX  /  CONCEPTS LEVEL 0',CONCEPTS,'A1','concepts',
    '旧分类体系：Level 0 为主学科，Level 1 为细分概念。19 个主学科按撤稿候选数量排序；标签可重叠。',CONCEPT_NOTE)
trend_slide('五个高数量主学科，沿发表年份展开', 'OPENALEX  /  CONCEPTS LEVEL 0 · OA 数量前五',TOP_CONCEPTS,'A1','concepts')
for root in TOP_CONCEPTS:
    detail_slide(root,'concepts','A1')

item=slide('Topics：用主主题，把每篇文献放到一条分类路径', '方法  /  DOMAIN → FIELD → SUBFIELD → TOPIC',
    '当前分类体系由模型分配；本研究使用 primary_topic，不把多个 topics 关联重复计入同一层。',
    '官方方法：https://help.openalex.org/data/topics/ ；缺失主主题单列，不补造类别。')
text_columns(item,[('四层路径','Domain：大领域\nField：学科\nSubfield：子学科\nTopic：具体研究主题'),
    ('本报告比较两层','先看全部 26 个 Field，\n再看高数量 Field 内的 Subfield。\nOA 与 RW 采用同一套节点。'),
    ('与 Concepts 不同','主主题每层只计一次。\nConcepts 可以同时关联多学科。\n两套分类并排研究，不混为同一分母。')])

distribution('Topics Field：数量与学科内比例并排看', 'OPENALEX  /  PRIMARY TOPIC → FIELD',FIELDS,'A1','topics',
    '26 个 Field 按 OA 撤稿候选数量排序；右侧沿用相同顺序，不是比例排行榜。')
trend_slide('从医学到工程：比例趋势并不只是数量曲线的缩放', 'OPENALEX  /  TOPICS FIELD · OA 数量前五',TOP_FIELDS,'A1','topics')
for root in FIELDS[:3]:
    detail_slide(root,'topics','A1')

item=slide('RW：先匹配原文，再进入共同分析范围', 'RW × OPENALEX  /  单位与时间范围',
    'RW 提供撤稿记录与原因；OpenAlex 补充主学科和全部发文分母。精确 DOI / PMID 匹配，不任意选择歧义候选。',
    '60,313 / 66,700 = 已匹配 RW 原论文占比；60,311 为对应不同 OA Work 数，单位不同，不混算匹配率。')
table(item,['阶段','数量','计数单位'],[
    ['RW 快照内去重原论文',f'{ANALYSIS["rw_originals"]:,}','RW 原论文'],
    ['成功匹配 OA 的原论文',f'{ANALYSIS["rw_linked_originals"]:,}',f'RW 原论文；{ANALYSIS["rw_linked_originals"]/ANALYSIS["rw_originals"]*100:.2f}%'],
    ['对应的不同 OA 记录',f'{ANALYSIS["rw_matched_distinct_works"]:,}','OA Work'],
    ['通过 OA 范围与共同事件截止',f'{ANALYSIS["rw_eligible_distinct_all_years"]:,}','不同 OA Work；全部发表年'],
    ['限定 2000–2025 发表队列',f'{ANALYSIS["rw_eligible_distinct_window"]:,}','后续 RW 学科与原因分析总体']],font_size=17)

distribution('匹配 RW：沿用同一组 Field，观察另一条记录线', 'RW × OPENALEX  /  PRIMARY TOPIC → FIELD',FIELDS,'C_D','topics',
    '分类与排列均沿用 OA；RW 分子为截止日前撤稿的匹配原文，分母仍是同口径 OA 全部发文。')
trend_slide('匹配 RW：同一组 Field 的发表年队列', 'RW × OPENALEX  /  TOPICS FIELD · 与 OA 相同五类',TOP_FIELDS,'C_D','topics')
for root in FIELDS[:3]:
    detail_slide(root,'topics','C_D')

item=slide('同一分母下，两条观察线仍然不完全重合', 'OA 与匹配 RW  /  FIELD 比例对照',
    '两库并非独立验证：OA 撤稿标记源自 RW。差异可能涉及覆盖、时间、匹配及记录身份，不在此作因果归因。',COMMON)
figure,axis=plt.subplots(figsize=(14.55,5.2))
figure.subplots_adjust(left=.25,right=.88,bottom=.12,top=.93)
for index,node in enumerate(TOP_FIELDS):
    values=[proportion(node,'A1'),proportion(node,'C_D')]
    axis.plot(values,[index,index],color='#C0CBCD',linewidth=3)
    axis.scatter(values,[index,index],color=[BLUE,RED],s=85,zorder=3)
    axis.text(values[0],index-.15,f'OA {values[0]:.3f}%',fontsize=12,color=BLUE)
    axis.text(values[1],index+.27,f'RW {values[1]:.3f}%',fontsize=12,color=RED)
axis.set_yticks(range(5),[label(node) for node in TOP_FIELDS])
axis.set_xlabel('学科内观测撤稿比例（%）')
axis.set_xlim(0,max(proportion(node) for node in TOP_FIELDS)*1.3)
axes_style(axis,True)
embed(item,save_figure(figure,'comparison'))
note_evidence(item,TOP_FIELDS,'A1','topics')
note_evidence(item,TOP_FIELDS,'C_D','topics')

item=slide('撤稿原因：实质标签、疑虑与程序信息必须分开', 'RW 匹配样本  /  原因族覆盖比例',
    f'分母为 {ANALYSIS["rw_eligible_distinct_window"]:,} 条匹配合格 Work；同篇标签先取并集，每个原因族只计一次。',
    '事件截至 2026-06-26；标签取 RW 2026-09-10 快照，不能视为截止时已知。多标签可重叠；原因族为项目归组。')
reason_order=sorted(ANALYSIS['reason_counts'],key=lambda key:-ANALYSIS['reason_counts'][key])
figure,axis=plt.subplots(figsize=(14.55,5.2))
figure.subplots_adjust(left=.25,right=.91,bottom=.10,top=.97)
counts=[ANALYSIS['reason_counts'][key] for key in reason_order]
values=[number/ANALYSIS['rw_eligible_distinct_window']*100 for number in counts]
axis.barh(range(len(values)),values,color=RED,height=.60)
axis.set_yticks(range(len(values)),[ANALYSIS['reason_names'][key] for key in reason_order],fontsize=11)
axis.set_xlabel('匹配撤稿样本内覆盖比例（%），不是全部发文撤稿比例')
axis.set_xlim(0,max(values)*1.26)
for index,(number,value) in enumerate(zip(counts,values)):
    axis.text(value+.4,index,f'{value:.1f}% · {number:,}',va='center',fontsize=10)
axes_style(axis,True)
embed(item,save_figure(figure,'reasons'))
item.notes_slide.notes_text_frame.text=json.dumps({'denominator':ANALYSIS['rw_eligible_distinct_window'],'counts':ANALYSIS['reason_counts'],'mapping':'rw-exact-label-families-v1','missing':ANALYSIS['reason_counts'].get('missing',0),'unmapped':ANALYSIS['reason_counts'].get('unmapped',0),'names':ANALYSIS['reason_names']},ensure_ascii=False)

item=slide('不同学科，记录的原因组合也不同', 'RW 匹配样本  /  FIELD × 原因族',
    '每一列的分母是该 Field 内匹配撤稿文献数，不是该 Field 全部发文数；展示总体覆盖数量前 8 个原因族。',
    '事件截至 2026-06-26；原因标签取 RW 2026-09-10 快照，非截止时已知。多标签可重叠，非学术不端裁定。')
reason_subset=reason_order[:8]
matrix=np.array([[100*ANALYSIS['field_reasons'][node['id']].get(key,0)/ANALYSIS['field_reason_denominators'][node['id']] for node in TOP_FIELDS] for key in reason_subset])
figure,axis=plt.subplots(figsize=(14.55,5.2))
figure.subplots_adjust(left=.25,right=.91,bottom=.25,top=.97)
image=axis.imshow(matrix,cmap='YlOrBr',vmin=0,vmax=100,aspect='auto')
axis.set_xticks(range(5),[label(node).replace('、','\n')+f'\nn={count(node,"C_D"):,}' for node in TOP_FIELDS],fontsize=11)
axis.set_yticks(range(8),[ANALYSIS['reason_names'][key] for key in reason_subset],fontsize=12)
axis.tick_params(length=0)
for row in range(8):
    for column in range(5):
        axis.text(column,row,f'{matrix[row,column]:.1f}%',ha='center',va='center',color='white' if matrix[row,column]>65 else INK,fontsize=13)
figure.colorbar(image,ax=axis,fraction=.03,pad=.03,label='该 Field 匹配撤稿样本覆盖（%）')
embed(item,save_figure(figure,'reason-heatmap'))
item.notes_slide.notes_text_frame.text=json.dumps({'denominators':ANALYSIS['field_reason_denominators'],'counts':ANALYSIS['field_reasons'],'rows':reason_subset,'columns':[node['id'] for node in TOP_FIELDS]},ensure_ascii=False)

item=slide('近期曲线下降，首先要排查观察窗口与标签覆盖', '解释边界  /  时间与分类缺失',
    '比例控制发文规模，但并未消除撤稿时滞、标签缺失、匹配选择或宽口径身份不确定性。',
    'Concepts 已停止更新；这里只分析快照已有标签。2026 年不完整，仅列附注，不并入主图。')
missing_rows=[]
for taxonomy,level,name in [('concepts',0,'Concepts L0'),('concepts',1,'Concepts L1'),('topics',1,'Topics Field'),('topics',2,'Topics Subfield')]:
    missing=[node for node in TAX[taxonomy]['nodes'] if node['level']==level and node.get('missing') and node['denominator'] is not None]
    missing_rows.append([name,f'{sum(count(node) for node in missing):,}',f'{sum(count(node,"C_D") for node in missing):,}',f'{sum(denominator(node) for node in missing):,}'])
table(item,['2000–2025 缺失类别','OA 撤稿候选缺失','匹配 RW 缺失','全部发文分类缺失'],missing_rows,height=2.8,font_size=16)
topics=TAX['topics']
text(item,f'2026 不完整队列（未入主图）：OA {topics["counts"]["A1"][27]:,} 条；匹配 RW {topics["counts"]["C_D"][27]:,} 条；发文 N={topics["denominator"][27]:,}。',.85,5.7,14.1,.7,18)
text(item,'不同发表年的随访长度不同。没有固定随访或调整模型，不能声称学科风险发生因果变化。',.85,6.65,14.1,.7,18,False,RED)

item=slide('答案不是一张“高风险学科榜”', '研究结论  /  数量、比例、原因各回答不同问题', caption='本研究为描述性观察；不能用于对个人、机构或整个学科作学术不端裁定。')
text_columns(item,[('数量：记录在哪里',f'医学主 Field 为 {count(medicine):,} 条 OA 候选。\nConcepts 计算机科学为 {count(TOP_CONCEPTS[0]):,} 条关联记录。\n两者分类机制不同。'),
    ('比例：放回发文基数',f'生化、遗传与分子生物学：{proportion(biochem):.3f}%。\n医学：{proportion(medicine):.3f}%。\n排序取决于分母，不仅取决于数量。'),
    ('原因：理解记录内容','区分错误、疑虑、可靠性问题与调查程序。\n匹配样本支持同类比较，但不能代表全部 RW 或未匹配文献。')])

item=slide('关键数值：保留 n / N，比例才可以核对', '附录  /  2000–2025 全期 · 两套分类分别统计',
    '完整年度数值与节点 ID 保存在对应幻灯片备注；未绘制的小基数单元保留真实值。',COMMON)
rows=[]
for node in TOP_CONCEPTS:
    rows.append(['Concepts L0',label(node),f'{count(node):,}',f'{denominator(node):,}',f'{proportion(node):.3f}%'])
for node in TOP_FIELDS:
    rows.append(['Topics Field',label(node),f'{count(node):,}',f'{denominator(node):,}',f'{proportion(node):.3f}%'])
grid=table(item,['体系','学科','OA 分子 n','发文分母 N','比例'],rows,font_size=13)
grid.columns[0].width=Inches(1.65)
grid.columns[1].width=Inches(4.4)
grid.columns[2].width=Inches(2.3)
grid.columns[3].width=Inches(3.1)
grid.columns[4].width=Inches(2.95)

item=slide('来源、版本与可复核的方法边界', '附录  /  EVIDENCE & METHODS',
    f'发布标识：{ANALYSIS["release_id"]}  ·  本专题制作：2026-09-16',
    '公开材料仅含聚合统计；无逐篇原始记录、作者名单或凭据。未更改现有网站、源快照和已发布聚合。')
text(item,'数据证据',.85,2.5,7,.5,23,True)
text(item,'public/data/snapshot/manifest.json、fields.json\n全类型发文分母与候选：宽口径既有统计\n匹配与原因：去重 RW 原文 → 不同 OA Work\n方法版本：original-first-independent-notices-v2\n原因族：rw-exact-label-families-v1',.85,3.2,7.1,3.4,17)
text(item,'分类与数据源说明',8.25,2.5,7,.5,23,True)
links=[('OpenAlex Works','https://help.openalex.org/data/works/'),('OpenAlex Topics','https://help.openalex.org/data/topics/'),('Concepts（旧体系）','https://help.openalex.org/data/concepts/'),('Crossref / Retraction Watch','https://www.crossref.org/documentation/retrieve-metadata/retraction-watch/'),('RW 原因字段释义','https://retractionwatch.com/retraction-watch-database-user-guide/retraction-watch-database-user-guide-appendix-b-reasons/')]
for index,(heading,url) in enumerate(links):
    shape=text(item,heading+' ↗',8.25,3.15+index*.68,7,.5,19,False,BLUE)
    shape.text_frame.paragraphs[0].runs[0].hyperlink.address=url
item.notes_slide.notes_text_frame.text=json.dumps({'manifest':DATA['manifest'],'analysis_checks':ANALYSIS['checks'],'sources':links},ensure_ascii=False)

assert len(DECK.slides)==30,len(DECK.slides)
path=OUT/'撤稿集中在哪里_学科专题_2000-2025.pptx'
DECK.save(path)
with (OUT/'chart-evidence.csv').open('w',newline='') as handle:
    writer=csv.DictWriter(handle,fieldnames=list(EVIDENCE[0]))
    writer.writeheader()
    writer.writerows(EVIDENCE)
(OUT/'slides.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2))
(OUT/'axis-audit.json').write_text(json.dumps(AXIS_AUDIT,ensure_ascii=False,indent=2))
print(json.dumps({'pptx':str(path),'slides':len(DECK.slides),'evidence_rows':len(EVIDENCE),'bytes':path.stat().st_size},ensure_ascii=False))
