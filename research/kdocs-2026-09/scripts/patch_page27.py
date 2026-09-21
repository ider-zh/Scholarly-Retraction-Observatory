import io
import json
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path('/tmp/retraction-kdocs-20260916')
original = OUT / '撤稿集中在哪里_学科专题_2000-2025.pptx'
presentation = Presentation(original)
page = presentation.slides[26]

def set_text(shape, value):
    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    for extra in list(frame.paragraphs)[1:]:
        frame._txBody.remove(extra._p)
    if paragraph.runs:
        paragraph.runs[0].text = value
        for extra in list(paragraph.runs)[1:]:
            paragraph._p.remove(extra._r)
    else:
        paragraph.text = value

set_text(page.shapes[1], '缺少标签的文献记录数，不是缺少学科类别')
set_text(page.shapes[2], 'L0 共有 19 个主学科，均有记录；14 指缺少 L0 标签的 OA 撤稿候选文献，不是缺少 14 个学科。')
set_text(page.shapes[4], '2000–2025 年发表队列，观察截至 2026-06-26。每列分母是该列全部文献数；不是学科类别数。')
set_text(page.shapes[8], 'L1 全部合格发文的标签缺失比例为 20.9147%，需要注意覆盖差异。\n有标签不等于分类准确；Concepts 保留既有标签，不重新分类。')
page.shapes[8].top = Inches(5.85)
page.shapes[8].height = Inches(.8)
page.shapes[8].text_frame.paragraphs[0].font.size = Pt(17)
set_text(page.shapes[9], '本页仅描述标签覆盖，未检验曲线下降的原因；不能据此归因于标签缺失或观察窗口。')
page.shapes[9].top = Inches(6.85)
page.shapes[9].text_frame.paragraphs[0].font.size = Pt(17)
table = page.shapes[7].table
headers = ['分类层级', 'OA 撤稿候选\n缺失文献数 / 总文献数', '匹配 RW\n缺失文献数 / 总文献数', 'OA 全部合格发文\n缺失文献数 / 总文献数']
taxonomies = {item['id']: item for item in json.loads((OUT/'published.json').read_text())['fields']['discipline_explorer']['taxonomies']}
rows = [headers]
for taxonomy, level, name in [('concepts',0,'Concepts L0'),('concepts',1,'Concepts L1'),('topics',1,'Topics Field'),('topics',2,'Topics Subfield')]:
    nodes = [node for node in taxonomies[taxonomy]['nodes'] if node['level']==level and node.get('missing')]
    values = [sum(sum(node['counts'][population][1:27]) for node in nodes) for population in ['A1','C_D']]
    values.append(sum(sum(node['denominator'][1:27]) for node in nodes))
    totals = [77414,59028,220305891]
    rows.append([name]+[f'{number:,} / {total:,}\n（{number/total*100:.4f}%）' for number,total in zip(values,totals)])
for index,row in enumerate(rows):
    table.rows[index].height = Inches(.65)
    for column,value in enumerate(row):
        cell = table.cell(index,column)
        cell.text = value
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = 'Noto Sans CJK SC'
            paragraph.font.size = Pt(14)
            paragraph.font.bold = index==0
            paragraph.font.color.rgb = RGBColor.from_string('182B35')
            paragraph.space_after = Pt(0)
            paragraph.space_before = Pt(0)
buffer = io.BytesIO()
presentation.save(buffer)
with ZipFile(buffer) as generated:
    replacement = generated.read('ppt/slides/slide27.xml')
path = OUT/'updated-deck.pptx'
with ZipFile(original) as source, ZipFile(path,'w') as target:
    for entry in source.infolist():
        target.writestr(entry,replacement if entry.filename=='ppt/slides/slide27.xml' else source.read(entry.filename))
with ZipFile(original) as before, ZipFile(path) as after:
    changed = [name for name in before.namelist() if before.read(name)!=after.read(name)]
assert changed==['ppt/slides/slide27.xml'],changed
print('Only slide27.xml changed; all other package entries identical.')
