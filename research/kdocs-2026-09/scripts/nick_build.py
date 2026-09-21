import os

os.environ['PPT_OUTPUT_DIR'] = '/tmp/retraction-nick-20260920'
source_code = open('/tmp/retraction-kdocs-20260916/build_deck.py').read().split("item=slide('撤稿集中在哪里？'")[0]
exec(compile(source_code, 'presentation-helpers', 'exec'))

import hashlib
import re
from docx import Document
from docx.shared import Cm, Pt as WordPt, RGBColor as WordRGB
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

EVENT = json.loads((OUT / 'analysis.json').read_text())
SUMMARY = EVENT['summary']
GROUPS = EVENT['groups']
TRANSLATIONS.update({'Law': '法学', 'Electrical and Electronic Engineering': '电气与电子工程'})
INPUT = BASE / 'event-year/retraction-event-year.pptx'
DECK = Presentation(INPUT)
original = list(DECK.slides)
assert len(original) == 38

def update(shape, value):
    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in list(paragraph.runs)[1:]:
            run._r.getparent().remove(run._r)
    else:
        paragraph.text = value
    for extra in list(frame.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)

def main_note(item, extra=''):
    item.notes_slide.notes_text_frame.text = json.dumps({'method': SUMMARY, 'explanation': extra}, ensure_ascii=False)

cover = slide('OpenAlex 撤稿的学科分布与年度变化', '学科撤稿观察 / 研究简报',
    '基于 RW—OpenAlex 匹配样本的撤稿年份观察，附 OpenAlex 发表年份分析',
    'OA 快照 2026-06-26；RW 快照 2026-09-10。正文不是 OA 撤稿标记总体，也不要求 OA 同时标记撤稿。')
text(cover, '哪些学科的撤稿记录更多？\n这些记录如何随年份变化？', .8, 2.85, 13.8, 1.7, 38, True)
line(cover, .8, 5.0, 14.2)
text(cover, '正文 / 撤稿年份视角', .8, 5.30, 6.7, .5, 24, True, BLUE)
text(cover, '58,345 条合格匹配文献\n2000—2025 年被记录撤稿', .8, 6.0, 6.7, 1.05, 23)
text(cover, '附录 / 仅 OpenAlex', 8.1, 5.30, 6.7, .5, 24, True, RED)
text(cover, '77,414 条 OA 撤稿标记候选\n2000—2025 年发表队列', 8.1, 6.0, 6.7, 1.05, 23)
main_note(cover)

methods = slide('研究问题、样本与比例：先确定分母', '正文 / RW—OPENALEX 匹配样本',
    'RW 提供撤稿记录与日期；OpenAlex 提供匹配 Work、学科标签和共同资格筛选。',
    'Work 是学术文献记录，含论文、书籍章节等；正文比例描述撤稿文献构成，不是学科撤稿风险。')
text_columns(methods, [
    ('怎样形成样本', '精确 DOI / PMID 匹配，按不同 OA Work 去重。\n60,311 条匹配 Work → 60,288 条合格记录 → 58,345 条在研究期内撤稿。\n采用主体库（core）全类型宽口径；不要求 OA 撤稿标记为真。'),
    ('怎样组织时间', '取最早可解析 RW 撤稿日期，统计 2000—2025 年。\n不限制原文发表于该窗口；含 516 条更早发表文献。\n部分日期存在不确定性，年度分布不等于逐篇核实的实际事件分布。'),
    ('怎样计算比例', '当年某学科撤稿数 n ÷ 当年全部合格匹配撤稿数 N × 100%。\nN 包含分类缺失，不是该学科全部发文。\n全期分布用全期总数作分母；前六按全期数量固定选取。')
], top=2.5)
text(methods, f"日期质量：{SUMMARY['date_uncertainty_tagged_works']:,} 条带日期不确定标记；另有 213 条 OA 发表日期晚于 RW 撤稿日期。均保留，不推断时滞。", .75, 7.02, 14.5, .45, 14, False, MUTED)
main_note(methods)

classification = slide('同一批文献，两套学科分类', '正文 / 分类方法',
    '先用 Concepts 延续前期研究，再用 Topics 比较学科分布与年度变化；不是独立数据验证。',
    'Concepts L1 为该标签全体文献，不与某一父标签取交集；Topics 使用 primary_topic 路径，缺失分类单列。')
table(classification, ['比较维度','Concepts：概念标签','Topics：主主题路径'], [
    ['层级','L0 主学科；L1 次级概念','Domain → Field → Subfield → Topic'],
    ['本报告研究层级','L0（一级）、L1（二级）','Field、Subfield'],
    ['一篇文献如何分类','可以关联多个概念标签','沿 primary_topic，每层只归一次'],
    ['能否相加','标签重叠，比例和可超过 100%','同层含缺失项后，比例和为 100%'],
    ['为什么两套都用','延续前期概念关联观察','观察单一主类分配后的变化'],
    ['如何比较','不将 L0/L1 与 Field/Subfield 视为一一对应','比较排序与趋势，不把主主题视为真值']
],font_size=17)
main_note(classification)

comparison = slide('一致的是交叉分布，变化的是归类与排序', '正文 / 两套体系结果比较',
    '同一组 58,345 条文献、同一撤稿年份口径；差异来自分类规则，不是样本换了一批。',
    'Concepts 标签关联不能直接解释为严格主学科归属；同名或近似学科在两套体系中也不保证覆盖一致。')
text_columns(comparison, [
    ('哪些观察一致', '医学、计算机、工程及生命科学相关分类，在两套体系中均有较多记录。\n分子生物学、癌症研究、人工智能等在细分类别中出现。\n数量变化与占比变化回答不同问题。'),
    ('哪些排序改变', 'Concepts L0：计算机科学 25,748，医学 21,405。\nTopics Field：医学 12,957，计算机科学 7,337。\n多标签关联转为主类分配后，首位由计算机科学变为医学。'),
    ('怎样解释峰值', 'Concepts L1 法学标签在 2011 年为 3,969；这些记录的主 Topics 分散在多个 Field。\n不能将其直接称为法学论文的真实撤稿高峰。\n分类不同与日期质量都需纳入解释。')
],top=2.5)
main_note(comparison, '2011 年 Law 局部核验：3969 条；主 Field 商业管理777、环境647、工程640、经济410、社会366、计算机303、其他及缺失826；主 Subfield 为 Law 3 条。未调整标签规则。')

conclusion = slide('回答研究问题：分清数量、份额与分类', '正文 / 结论',
    '本报告说明匹配样本中的学科分布及其年度变化，不估计学科的真实不端发生率。',
    '附录另用 OA 全部合格发文作分母；正文年度构成比例与附录学科内撤稿标记比例不可混用。')
text_columns(conclusion, [
    ('学科分布', 'Concepts 延续前期研究：L0 计算机科学关联记录最多；L1 法学标签全期关联量最多。\nTopics 主 Field 中医学最多；Subfield 中分子生物学最多。'),
    ('年度变化', '整个匹配样本中，2023 年记录数为 12,873，是研究期内最高年份。\n部分分类在 2010—2011 年也有明显峰值。\n年度占比下降，不一定意味着数量减少。'),
    ('研究边界', '结论适用于合格 RW—OA 匹配样本，不外推到全部 OA 或未匹配 RW。\n保留日期不确定记录与多标签。\n附录提供不依赖 RW 的 OA 发表年份观察。')
],top=2.5)
main_note(conclusion)

original_ids = list(DECK.slides._sldIdLst)
order = [38,39,40,30,31,32,33,34,35,41,42,4,5,12,13]
for index, identifier in enumerate(original_ids):
    DECK.slides._sldIdLst.remove(identifier)
    if index not in order:
        DECK.part.drop_rel(identifier.rId)
for index in order:
    DECK.slides._sldIdLst.append(original_ids[index])
titles = {4:'Concepts 一级学科：撤稿记录的总体分布',5:'Concepts 一级学科：前六类的年度数量与占比',6:'Concepts 二级概念：前六个标签的年度变化',7:'Topics Field：主学科的撤稿记录分布',8:'Topics Field：前六类的年度数量与占比',9:'Topics Subfield：前六类的年度变化'}
for number,item in enumerate(DECK.slides,1):
    if number in titles:
        update(item.shapes[1],titles[number])
    if 4 <= number <= 9:
        update(item.shapes[0],'正文 / RW—OPENALEX 匹配样本 / 2000—2025 撤稿记录年')
    if number >= 12:
        update(item.shapes[0], '附录 / 仅 OPENALEX / 2000—2025 发表年份视角')
        original_caption = item.shapes[2].text
        update(item.shapes[2], '仅 OA：n 为带撤稿标记的合格文献；N 为同分类、同发表年全部合格 OA 发文。与正文匹配样本不同。')
    for shape in item.shapes:
        if not shape.has_text_frame:
            continue
        if shape.left > Inches(14.5) and shape.top > Inches(8):
            update(shape,f'{number:02d}')
        elif shape.top > Inches(8.2):
            update(shape, '正文 / RW—OA 匹配样本 58,345 · OA 2026-06-26 / RW 2026-09-10' if number <= 11 else
                '附录 / OA 撤稿标记 77,414 / 合格发文 220,305,891 · 发表年 2000—2025 · 快照 2026-06-26')
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace('Concepts 保留既有标签，不另设分数门槛；多标签重叠，L1 非父子交集。','Concepts 多标签可重叠；L1 按标签全体文献计数，不与父标签取交集。')
    if number >= 12:
        item.notes_slide.notes_text_frame.text += '\n附录独立口径：仅 OA core 全类型宽口径，2000—2025 发表队列，is_retracted 分子，同学科全部合格 OA 发文分母；不使用 RW 匹配或撤稿日期。'
    notes = item.notes_slide.notes_text_frame
    notes.text = re.sub(r'"page":\s*\d+', f'"page": {number}', notes.text)
DECK.core_properties.title = 'OpenAlex 撤稿的学科分布与年度变化'
DECK.core_properties.subject = '正文 RW—OpenAlex 匹配撤稿年份样本；附录仅 OpenAlex 发表年份样本'
DECK.save(OUT/'学科撤稿研究_尼克简报_20260920.pptx')

document = Document()
section = document.sections[0]
section.page_height,section.page_width = Cm(29.7),Cm(21)
section.top_margin,section.bottom_margin = Cm(1.8),Cm(1.8)
section.left_margin,section.right_margin = Cm(2),Cm(2)
for name in ['Normal','Heading 1','Heading 2','Title']:
    style=document.styles[name]
    style.font.name='Noto Sans CJK SC'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
document.styles['Normal'].font.size=WordPt(10.5)
document.styles['Normal'].paragraph_format.space_after=WordPt(6)
document.styles['Normal'].paragraph_format.line_spacing=1.15
document.styles['Heading 1'].font.size=WordPt(16)
document.styles['Heading 1'].font.color.rgb=WordRGB.from_string('176B87')
document.styles['Heading 2'].font.size=WordPt(12)
document.styles['Title'].font.size=WordPt(20)
header=section.header.paragraphs[0]
header.text='学科撤稿研究 | RW—OpenAlex 匹配样本 | 2026-09-20'
header.style=document.styles['Normal']
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页 / 正文按撤稿记录年；OA 发表年份结果见配套 PPT 附录')

def para(value):
    document.add_paragraph(value)

def word_table(headers,rows):
    grid=document.add_table(rows=1,cols=len(headers));grid.style='Light Shading Accent 1'
    for cell,value in zip(grid.rows[0].cells,headers):cell.text=value
    for row in rows:
        for cell,value in zip(grid.add_row().cells,row):cell.text=str(value)
    for row in grid.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=WordPt(2)
                for run in paragraph.runs:run.font.size=WordPt(9.5)

def group_rows(group):
    return [[label(node),f"{node['total']:,}",f"{node['total']/58345*100:.2f}%"] for node in GROUPS[group]['nodes'][:6]]

document.add_heading('OpenAlex 撤稿的学科分布与年度变化',0)
para('基于 RW—OpenAlex 匹配样本的撤稿年份观察，附 OpenAlex 发表年份分析')
document.add_heading('一、研究问题、数据和分类体系',1)
para('本研究回答：哪些学科关联了更多撤稿记录？各学科的撤稿数量及其在年度撤稿样本中的占比如何变化？先用 Concepts 延续前期研究，再用 Topics 比较分类规则改变后的分布与年度变化。两套体系作用于同一批文献，不是独立样本验证。')
para('数据版本为已验收的 OpenAlex 2026-06-26 快照和 Retraction Watch（RW）2026-09-10 快照，不称为截至撰写日的最新版本。RW 提供撤稿记录和日期，OpenAlex 提供匹配 Work、学科标签及共同资格筛选。按精确 DOI/PMID 匹配，不任意选择歧义候选；同一 OA Work 只计一次。正文不要求 OA 同时具有撤稿标记，因此称为“匹配样本”，而不是两库撤稿标记的严格交集。')
para('60,311 条匹配 Work 中，60,288 条通过共同资格条件，最终 58,345 条的最早可解析 RW 撤稿日期处于 2000—2025 年。使用 core 全类型宽口径，排除有明确证据的独立通知及不合格日期记录，保留身份未决项；不限定原文发表于研究时间窗，样本包含 516 条发表于 2000 年前的文献。')
word_table(['项目','Concepts','Topics'],[['使用层级','L0 主学科、L1 次级概念','Field、Subfield'],['同篇计数','多标签，学科间可能重复','沿 primary_topic，每层一次'],['比例合计','可以超过 100%','同层含缺失项后为 100%']])
para('年度学科占比＝当年该学科的合格匹配撤稿文献数 n／当年全部合格匹配撤稿文献数 N×100%。N 包含分类缺失；Subfield 不改用所属 Field 为分母。全期占比用全期总 n／58,345，而不是平均年度比例。前六按全期数量固定选取；N<100 用空心点提醒，N=0 留空，N>0 且 n=0 保留真实零值。')
para(f"日期限制：{SUMMARY['date_uncertainty_tagged_works']:,} 条文献带有“文章或通知日期不明”标签，仍按 RW 所填日期统计，不将其视为逐篇核实的实际撤稿事件日期。另有 213 条 OA 发表日期晚于 RW 撤稿日期，保留并披露，不推断撤稿时滞。")

document.add_page_break()
document.add_heading('二、Concepts 体系下的结果',1)
para('Concepts 分析的是文献与概念标签的关联。L0 展示全部 19 个主学科；L1 为次级概念，各标签按自身全体文献计数，不与某个父标签取交集，不能将子项相加推算父项。下表均为全期总量及其占全部 58,345 条匹配撤稿文献的比例。')
document.add_heading('1. 一级学科：计算机科学关联记录最多',2)
word_table(['L0 主学科','撤稿记录数','全期占比'],group_rows('concepts_l0'))
para('计算机科学、医学、生物学位居前三，但这些分类相互重叠；例如一篇文献可以同时进入医学和生物学。表中比例不是该学科全部发文被撤稿的概率。')
document.add_heading('2. 二级概念：数量与年度变化',2)
word_table(['L1 标签','撤稿记录数','全期占比'],group_rows('concepts_l1'))
para('法学标签全期关联 13,214 条，2010 年为 2,322 条、2011 年为 3,969 条；2011 年占当年 4,684 条匹配撤稿文献的 84.74%。这一结果必须称为“Law 标签关联记录”，不能直接写成“法学论文撤稿高峰”。其他高数量标签包括内科学、人工智能、生物化学、遗传学和癌症研究。')
para('局部交叉检查显示，2011 年带 Law 标签的记录在主 Topics 中分散于商科、环境科学、工程等多个 Field；只有 3 条的主 Subfield 为 Law。两套体系不同的归类方式，是解读这种差异的必要背景。年度峰值也需考虑会议来源集中及日期不确定性，不能据图作学科不端判断。')

document.add_page_break()
document.add_heading('三、Topics 体系下的结果',1)
para('Topics 的层级为 Domain → Field → Subfield → Topic。按每篇文献的 primary_topic 归入对应路径，每一层只计一次。全期 Field 和 Subfield 均有 393 条分类缺失，保留在共同分母中；分类缺失不解释为撤稿数为零。')
document.add_heading('1. Field：医学居首，排序不同于 Concepts',2)
word_table(['Field','撤稿记录数','全期占比'],group_rows('topics_field'))
para('医学占 22.21%，其后是生物化学、遗传与分子生物学，工程学及计算机科学。Concepts 中计算机科学居首，而主 Topic Field 中医学居首；这是同一匹配样本在多标签关联和单一主类分配下呈现的不同画像。')
document.add_heading('2. Subfield 与年度变化',2)
word_table(['Subfield','撤稿记录数','全期占比'],group_rows('topics_subfield'))
para('分子生物学、癌症研究与人工智能位列所选 Subfield 前三。这里的比例仍除以全部匹配撤稿文献，不是除以对应 Field 的撤稿总量。配套 PPT 展示这六类逐年的数量与年度占比，以相同名单对照两种指标。')
para('全样本年度记录数在 2023 年达到研究期最高的 12,873 条；2010 和 2011 年分别为 2,859 和 4,684 条。单个学科的数量峰值不必与全样本一致。若其数量增长慢于其他学科，占比仍可能下降，因此数量和占比需要同时阅读，不能将占比下降直接理解为撤稿减少。')

document.add_page_break()
document.add_heading('四、两套体系结果的比较与分析',1)
para('一致之处在于：医学、计算机、工程和生命科学相关类别在两套体系中均具有较多记录；细分类别中可见癌症研究、人工智能等研究方向。数量描述记录规模，年度占比描述当年撤稿样本的学科构成，两者始终回答不同的问题。')
para('变化之处在于：Concepts 的计算机科学关联 25,748 条，医学关联 21,405 条；Topics 主 Field 中医学为 12,957 条，计算机科学为 7,337 条。不能把这些差额当作漏数，更不能把两套体系的计数相加。Concepts 的 L0/L1 与 Topics 的 Field/Subfield 不是一一对应层级，同名类别也未必覆盖同一组文献。')
para('对 2011 年 Law 标签峰值的局部检验进一步说明这一点：3,969 条记录中，主 Field 商业管理 777 条、环境科学 647 条、工程 640 条、经济 410 条、社会科学 366 条、计算机科学 303 条，其他及缺失 826 条。它们并没有消失，而是在另一套分类下被分配到多个位置。Topics 因而用于交叉比较，而不是判定 Concepts 对错的真值。')
document.add_heading('五、结论',1)
para('对研究需求的回答是：本次匹配样本的 Concepts 一级学科以计算机科学关联量最多，二级标签以 Law 关联量最多；改按主 Topics 分类后，Field 以医学最多，Subfield 以分子生物学最多。排名必须与分类定义一并解释，不能作为“高风险学科榜”。')
para('时间方面，按 RW 所填撤稿日期观察，2023 年是全样本记录最多的年份，部分学科标签在 2010—2011 年也有明显峰值。年度分布受到记录覆盖、批量事件、匹配及日期质量的共同限制，本报告未估计这些因素的因果贡献，不将峰值视为真实不端发生率的变化。')
document.add_heading('附录口径与证据',2)
para('配套 PPT 第 12—15 页另列完全基于 OA 的 2000—2025 年发表队列：合格发文 220,305,891 条，撤稿标记候选 77,414 条。该处比例＝同学科撤稿标记文献／同学科全部合格 OA 发文，展示 Concepts L0 和 Topics Field 的总体分布及发表年份趋势。不使用 RW 匹配条件或撤稿日期；近期发表队列观察期较短。正文与附录不能混用样本、分母或时间轴。')
para('证据：项目发布标识 oa-2026-06-26-e5c7695e3e2a；全类型规则 original-first-independent-notices-v2。正文逐分类、逐年 n/N 见配套 PPT 备注及 all-subjects.csv；附录依据既有公开 fields 聚合与 manifest 校验。Concepts 保留既有标签，未重新分类；本次不修改原始快照、旧报告或公开聚合。')
document.save(OUT/'学科撤稿研究_综述_20260920.docx')
assert len(DECK.slides)==15
for new_index,old_index in [(3,30),(4,31),(5,32),(6,33),(7,34),(8,35),(11,4),(12,5),(13,12),(14,13)]:
    images=lambda item:[hashlib.sha256(shape.image.blob).hexdigest() for shape in item.shapes if shape.shape_type==13]
    assert images(DECK.slides[new_index])==images(original[old_index])
print('15 slides; all 10 inherited chart images unchanged; Word written with four intended pages.')
